# -*- coding: utf-8 -*-
from __future__ import annotations

import os
from datetime import datetime
from pathlib import Path

import tkinter as tk
from tkinter import filedialog, messagebox, ttk

from openpyxl import Workbook


def collect_subdir_names(root: str) -> list[str]:
    names: list[str] = []
    for current, subdirs, _files in os.walk(root):
        for name in subdirs:
            names.append(name)
    names.sort(key=lambda x: x.lower())
    return names


def build_default_output_path(root: str) -> str:
    base = Path(root)
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    return str(base / f"subdir_names_{stamp}.xlsx")


def export_names_to_excel(names: list[str], output_path: str) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "目录名称"
    ws.cell(row=1, column=1, value="子目录名称")
    for idx, name in enumerate(names, start=2):
        ws.cell(row=idx, column=1, value=name)
    ws.column_dimensions["A"].width = 60
    Path(output_path).parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)


class App(tk.Tk):
    def __init__(self) -> None:
        super().__init__()
        self.title("子目录名称导出到Excel")
        self.geometry("760x420")
        self.minsize(700, 360)

        self.root_var = tk.StringVar()
        self.output_var = tk.StringVar()
        self.count_var = tk.StringVar(value="-")

        self._build_ui()

    def _build_ui(self) -> None:
        top = ttk.Frame(self)
        top.pack(fill="x", padx=12, pady=12)

        ttk.Label(top, text="根目录:").grid(row=0, column=0, sticky="w")
        ttk.Entry(top, textvariable=self.root_var, width=78).grid(row=0, column=1, padx=8, sticky="we")
        ttk.Button(top, text="选择...", command=self.choose_root).grid(row=0, column=2, sticky="e")

        ttk.Label(top, text="输出Excel:").grid(row=1, column=0, sticky="w", pady=(10, 0))
        ttk.Entry(top, textvariable=self.output_var, width=78).grid(row=1, column=1, padx=8, sticky="we", pady=(10, 0))
        ttk.Button(top, text="选择...", command=self.choose_output).grid(row=1, column=2, sticky="e", pady=(10, 0))

        ttk.Label(top, text="子目录数量:").grid(row=2, column=0, sticky="w", pady=(10, 0))
        ttk.Label(top, textvariable=self.count_var).grid(row=2, column=1, sticky="w", pady=(10, 0))

        btns = ttk.Frame(top)
        btns.grid(row=3, column=1, sticky="w", pady=(16, 0))
        ttk.Button(btns, text="扫描统计", command=self.scan_names).pack(side="left")
        ttk.Button(btns, text="导出Excel", command=self.export_excel).pack(side="left", padx=10)

        top.columnconfigure(1, weight=1)

        log_frame = ttk.LabelFrame(self, text="日志")
        log_frame.pack(fill="both", expand=True, padx=12, pady=(0, 12))

        self.log_text = tk.Text(log_frame, wrap="word")
        self.log_text.pack(fill="both", expand=True, padx=8, pady=8)
        self.log_text.configure(state="disabled")

    def _log(self, text: str) -> None:
        self.log_text.configure(state="normal")
        self.log_text.insert("end", text + "\n")
        self.log_text.see("end")
        self.log_text.configure(state="disabled")

    def choose_root(self) -> None:
        folder = filedialog.askdirectory(title="选择根目录")
        if not folder:
            return
        self.root_var.set(folder)
        if not self.output_var.get().strip():
            self.output_var.set(build_default_output_path(folder))
        self._log(f"已选择根目录: {folder}")

    def choose_output(self) -> None:
        initial = self.output_var.get().strip() or "subdir_names.xlsx"
        path = filedialog.asksaveasfilename(
            title="选择输出Excel",
            defaultextension=".xlsx",
            filetypes=[("Excel 文件", "*.xlsx")],
            initialfile=os.path.basename(initial),
        )
        if not path:
            return
        self.output_var.set(path)
        self._log(f"已选择输出文件: {path}")

    def scan_names(self) -> None:
        root = self.root_var.get().strip()
        if not root or not os.path.isdir(root):
            messagebox.showerror("错误", "根目录无效，请重新选择。")
            return
        names = collect_subdir_names(root)
        self.count_var.set(str(len(names)))
        self._log(f"扫描完成：共读取到 {len(names)} 个子目录名称。")

    def export_excel(self) -> None:
        root = self.root_var.get().strip()
        if not root or not os.path.isdir(root):
            messagebox.showerror("错误", "根目录无效，请重新选择。")
            return

        output_path = self.output_var.get().strip()
        if not output_path:
            output_path = build_default_output_path(root)
            self.output_var.set(output_path)

        names = collect_subdir_names(root)
        self.count_var.set(str(len(names)))
        export_names_to_excel(names, output_path)
        self._log(f"导出完成：{output_path} | 名称数量={len(names)}")
        messagebox.showinfo("完成", f"已导出到 Excel：\n{output_path}\n\n共 {len(names)} 个名称。")


if __name__ == "__main__":
    app = App()
    app.mainloop()
