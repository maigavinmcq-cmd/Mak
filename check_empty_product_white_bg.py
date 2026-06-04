# -*- coding: utf-8 -*-
from __future__ import annotations

import argparse
import sys
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
except ImportError:  # pragma: no cover - exercised by real CLI environments.
    Workbook = None
    Font = None
    PatternFill = None


DEFAULT_ROOT = r"\\192.168.1.6\004.短视频运营中心\01.产品信息\SG"
TARGET_FOLDER_NAME = "01.产品白底图"
IMAGE_EXTENSIONS = {
    ".jpg",
    ".jpeg",
    ".png",
    ".webp",
    ".bmp",
    ".gif",
    ".tif",
    ".tiff",
    ".heic",
    ".heif",
}


@dataclass(frozen=True)
class AuditRow:
    status: str
    folder_path: str
    product_dir: str
    image_count: int
    image_files: list[str]
    note: str


def is_image_file(path: Path) -> bool:
    return path.is_file() and path.suffix.lower() in IMAGE_EXTENSIONS


def scan_white_bg_folders(root: str | Path) -> list[AuditRow]:
    root_path = Path(root)
    if not root_path.is_dir():
        raise FileNotFoundError(f"根目录不存在或无法访问: {root_path}")

    rows: list[AuditRow] = []
    for product_dir in root_path.rglob("*"):
        if not product_dir.is_dir() or product_dir.name == TARGET_FOLDER_NAME:
            continue
        if any(parent.name == TARGET_FOLDER_NAME for parent in product_dir.parents):
            continue

        folder = product_dir / TARGET_FOLDER_NAME
        if not folder.is_dir():
            rows.append(
                AuditRow(
                    status="无白底图",
                    folder_path="",
                    product_dir=str(product_dir),
                    image_count=0,
                    image_files=[],
                    note="缺少 01.产品白底图 文件夹",
                )
            )
            continue

        image_files = sorted(
            (child.name for child in folder.iterdir() if is_image_file(child)),
            key=str.lower,
        )
        image_count = len(image_files)
        is_empty = image_count == 0
        rows.append(
            AuditRow(
                status="无图片" if is_empty else "有图片",
                folder_path=str(folder),
                product_dir=str(product_dir),
                image_count=image_count,
                image_files=image_files,
                note="需要补充白底图" if is_empty else "",
            )
        )

    rows.sort(key=lambda row: (row.product_dir.lower(), row.folder_path.lower()))
    return rows


def build_default_output_path(root: str | Path) -> Path:
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    return Path(root) / f"产品白底图检查结果_{stamp}.xlsx"


def export_to_excel(rows: list[AuditRow], output_path: str | Path) -> None:
    if Workbook is None:
        raise RuntimeError("缺少 openpyxl，请先安装：pip install openpyxl")

    output = Path(output_path)
    output.parent.mkdir(parents=True, exist_ok=True)

    wb = Workbook()
    ws = wb.active
    ws.title = "白底图检查"

    headers = ["序号", "状态", "白底图文件夹路径", "上级产品目录", "直属图片数量", "直属图片文件名", "备注"]
    ws.append(headers)

    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    empty_fill = PatternFill("solid", fgColor="FFF2CC")
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font

    for index, row in enumerate(rows, start=1):
        ws.append(
            [
                index,
                row.status,
                row.folder_path,
                row.product_dir,
                row.image_count,
                "\n".join(row.image_files),
                row.note,
            ]
        )
        if row.status in {"无图片", "无白底图"}:
            for cell in ws[ws.max_row]:
                cell.fill = empty_fill

    widths = {
        "A": 8,
        "B": 12,
        "C": 90,
        "D": 70,
        "E": 14,
        "F": 45,
        "G": 22,
    }
    for column, width in widths.items():
        ws.column_dimensions[column].width = width

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    wb.save(output)


def parse_args(argv: list[str]) -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="检查产品白底图文件夹是否缺少直属图片文件。")
    parser.add_argument("--root", default=DEFAULT_ROOT, help="要扫描的根目录，默认扫描 SG 产品信息目录。")
    parser.add_argument("--output", help="输出 Excel 路径，默认输出到根目录下并带时间戳。")
    return parser.parse_args(argv)


def main(argv: list[str] | None = None) -> int:
    args = parse_args(argv or sys.argv[1:])
    root = Path(args.root)
    output = Path(args.output) if args.output else build_default_output_path(root)

    try:
        rows = scan_white_bg_folders(root)
        export_to_excel(rows, output)
    except Exception as exc:
        print(f"检查失败: {exc}", file=sys.stderr)
        return 1

    no_image_count = sum(1 for row in rows if row.status == "无图片")
    no_folder_count = sum(1 for row in rows if row.status == "无白底图")
    print(f"扫描完成：共输出 {len(rows)} 条目录检查结果。")
    print(f"无直属图片：{no_image_count} 个。")
    print(f"无白底图文件夹：{no_folder_count} 个。")
    print(f"Excel 已输出：{output}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
