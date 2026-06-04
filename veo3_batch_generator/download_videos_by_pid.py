from __future__ import annotations

import argparse
import re
import time
import urllib.error
import urllib.parse
import urllib.request
from pathlib import Path

from openpyxl import load_workbook


DEFAULT_FILES = [
    Path("veo3_batch_generator/outputs/result_excel/Veo3视频生成结果_20260511_030251.xlsx"),
    Path("veo3_batch_generator/outputs/result_excel/Veo3视频生成结果_20260511_042113.xlsx"),
]


def safe_name(value: object, fallback: str) -> str:
    text = str(value or "").strip() or fallback
    return re.sub(r'[<>:"/\\|?*\x00-\x1f]', "_", text)


def find_column(headers: list[object], name: str) -> int:
    for idx, header in enumerate(headers):
        if str(header or "").strip() == name:
            return idx
    raise ValueError(f"Missing required column: {name}")


def video_id_from_url(url: str) -> str:
    parsed = urllib.parse.urlparse(url)
    slug = parsed.path.rstrip("/").split("/")[-1]
    return safe_name(slug, "video")


def iter_video_rows(workbook_path: Path):
    workbook = load_workbook(workbook_path, read_only=True, data_only=False)
    for sheet in workbook.worksheets:
        rows = sheet.iter_rows(values_only=True)
        try:
            headers = list(next(rows))
        except StopIteration:
            continue

        pid_col = find_column(headers, "PID")
        link_col = find_column(headers, "视频链接")

        for row_number, row in enumerate(rows, start=2):
            pid = row[pid_col] if pid_col < len(row) else None
            link = row[link_col] if link_col < len(row) else None
            if not isinstance(link, str) or not link.startswith(("http://", "https://")):
                continue
            yield {
                "pid": safe_name(pid, "UNKNOWN_PID"),
                "url": link,
                "row": row_number,
                "sheet": sheet.title,
                "source": workbook_path.stem,
                "video_id": video_id_from_url(link),
            }


def unique_path(path: Path) -> Path:
    if not path.exists():
        return path
    stem = path.stem
    suffix = path.suffix
    for idx in range(2, 10_000):
        candidate = path.with_name(f"{stem}_{idx}{suffix}")
        if not candidate.exists():
            return candidate
    raise RuntimeError(f"Unable to create unique filename for {path}")


def download(url: str, output_path: Path, retries: int, timeout: int) -> None:
    tmp_path = output_path.with_suffix(output_path.suffix + ".part")
    request = urllib.request.Request(
        url,
        headers={"User-Agent": "Mozilla/5.0"},
    )

    for attempt in range(1, retries + 1):
        try:
            with urllib.request.urlopen(request, timeout=timeout) as response:
                output_path.parent.mkdir(parents=True, exist_ok=True)
                with tmp_path.open("wb") as file:
                    while True:
                        chunk = response.read(1024 * 1024)
                        if not chunk:
                            break
                        file.write(chunk)
            tmp_path.replace(output_path)
            return
        except (urllib.error.URLError, TimeoutError, OSError):
            if tmp_path.exists():
                tmp_path.unlink()
            if attempt == retries:
                raise
            time.sleep(min(2 * attempt, 10))


def main() -> int:
    parser = argparse.ArgumentParser(description="Download Veo3 Excel video links grouped by PID.")
    parser.add_argument("workbooks", nargs="*", type=Path, default=DEFAULT_FILES)
    parser.add_argument(
        "--output-dir",
        type=Path,
        default=Path.home() / "Downloads" / "Veo3视频按PID",
    )
    parser.add_argument("--retries", type=int, default=3)
    parser.add_argument("--timeout", type=int, default=120)
    parser.add_argument("--limit", type=int, default=0, help="Download only the first N videos; 0 means all.")
    args = parser.parse_args()

    records = []
    seen_urls = set()
    for workbook_path in args.workbooks:
        workbook_path = workbook_path.resolve()
        for record in iter_video_rows(workbook_path):
            if record["url"] in seen_urls:
                continue
            seen_urls.add(record["url"])
            records.append(record)

    if args.limit:
        records = records[: args.limit]

    print(f"Found {len(records)} unique video links.")
    print(f"Output directory: {args.output_dir.resolve()}")

    downloaded = 0
    skipped = 0
    failed = []
    for index, record in enumerate(records, start=1):
        pid_dir = args.output_dir / record["pid"]
        filename = f"{record['source']}_row{record['row']}_{record['video_id']}.mp4"
        output_path = pid_dir / filename
        if output_path.exists() and output_path.stat().st_size > 0:
            skipped += 1
            print(f"[{index}/{len(records)}] skip {output_path}")
            continue

        output_path = unique_path(output_path)
        try:
            print(f"[{index}/{len(records)}] download PID={record['pid']} row={record['row']}")
            download(record["url"], output_path, retries=args.retries, timeout=args.timeout)
            downloaded += 1
        except Exception as exc:  # noqa: BLE001 - keep the batch moving and report all failures.
            failed.append((record, str(exc)))
            print(f"  FAILED: {exc}")

    print(f"Done. Downloaded={downloaded}, skipped={skipped}, failed={len(failed)}")
    if failed:
        print("Failed rows:")
        for record, error in failed:
            print(f"- {record['source']} row {record['row']} PID={record['pid']}: {error}")
        return 1
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
