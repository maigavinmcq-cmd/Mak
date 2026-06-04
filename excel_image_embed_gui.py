import hashlib
import io
import os
import shutil
import threading
import uuid
from pathlib import Path
from concurrent.futures import ThreadPoolExecutor, as_completed

import requests
import tkinter as tk
from PIL import Image as PILImage
from openpyxl import load_workbook
from openpyxl.drawing.image import Image as XLImage
from tkinter import filedialog, messagebox, ttk


def normalize_col(value: str) -> str:
    value = (value or "").strip().upper()
    if not value.isalpha():
        raise ValueError("列名必须是字母，例如 J 或 K")
    return value


class ExcelImageEmbedApp:
    def __init__(self, root: tk.Tk):
        self.root = root
        self.root.title("Excel 图片插入工具")
        self.root.geometry("920x700")

        self.worker = None
        self.stop_flag = False

        self.file_var = tk.StringVar()
        self.sheet_var = tk.StringVar()
        self.url_col_var = tk.StringVar(value="J")
        self.target_col_var = tk.StringVar(value="K")
        self.start_row_var = tk.StringVar(value="2")
        self.max_width_var = tk.StringVar(value="90")
        self.max_height_var = tk.StringVar(value="90")
        self.timeout_var = tk.StringVar(value="30")
        self.concurrent_var = tk.StringVar(value="8")
        self.copy_mode_var = tk.BooleanVar(value=True)
        self.backup_var = tk.BooleanVar(value=True)
        self.rebuild_var = tk.BooleanVar(value=False)
        self.force_redownload_var = tk.BooleanVar(value=False)

        self.cache_dir = Path.cwd() / "_tmp_excel_image_cache"
        self.cache_dir.mkdir(parents=True, exist_ok=True)

        self._build_ui()

    def _build_ui(self):
        frm_top = ttk.Frame(self.root, padding=12)
        frm_top.pack(fill="x")

        ttk.Label(frm_top, text="Excel 文件").grid(row=0, column=0, sticky="w")
        ttk.Entry(frm_top, textvariable=self.file_var, width=90).grid(row=0, column=1, sticky="ew", padx=8)
        ttk.Button(frm_top, text="选择文件", command=self.choose_file).grid(row=0, column=2, padx=4)
        ttk.Button(frm_top, text="读取工作表", command=self.load_sheets).grid(row=0, column=3, padx=4)
        frm_top.columnconfigure(1, weight=1)

        frm_cfg = ttk.LabelFrame(self.root, text="处理参数", padding=12)
        frm_cfg.pack(fill="x", padx=12, pady=(0, 8))

        ttk.Label(frm_cfg, text="工作表").grid(row=0, column=0, sticky="w")
        self.sheet_combo = ttk.Combobox(frm_cfg, textvariable=self.sheet_var, state="readonly", width=28)
        self.sheet_combo.grid(row=0, column=1, sticky="w", padx=(6, 18))

        ttk.Label(frm_cfg, text="图片链接列").grid(row=0, column=2, sticky="w")
        ttk.Entry(frm_cfg, textvariable=self.url_col_var, width=8).grid(row=0, column=3, sticky="w", padx=(6, 18))

        ttk.Label(frm_cfg, text="插入列").grid(row=0, column=4, sticky="w")
        ttk.Entry(frm_cfg, textvariable=self.target_col_var, width=8).grid(row=0, column=5, sticky="w", padx=(6, 18))

        ttk.Label(frm_cfg, text="起始行").grid(row=0, column=6, sticky="w")
        ttk.Entry(frm_cfg, textvariable=self.start_row_var, width=8).grid(row=0, column=7, sticky="w")

        ttk.Label(frm_cfg, text="最大宽").grid(row=1, column=0, sticky="w", pady=(10, 0))
        ttk.Entry(frm_cfg, textvariable=self.max_width_var, width=8).grid(row=1, column=1, sticky="w", padx=(6, 18), pady=(10, 0))

        ttk.Label(frm_cfg, text="最大高").grid(row=1, column=2, sticky="w", pady=(10, 0))
        ttk.Entry(frm_cfg, textvariable=self.max_height_var, width=8).grid(row=1, column=3, sticky="w", padx=(6, 18), pady=(10, 0))

        ttk.Label(frm_cfg, text="超时秒").grid(row=1, column=4, sticky="w", pady=(10, 0))
        ttk.Entry(frm_cfg, textvariable=self.timeout_var, width=8).grid(row=1, column=5, sticky="w", padx=(6, 18), pady=(10, 0))

        ttk.Label(frm_cfg, text="并发下载数").grid(row=2, column=0, sticky="w", pady=(10, 0))
        ttk.Entry(frm_cfg, textvariable=self.concurrent_var, width=8).grid(row=2, column=1, sticky="w", padx=(6, 18), pady=(10, 0))
        ttk.Checkbutton(frm_cfg, text="另存为 _with_images 副本", variable=self.copy_mode_var).grid(row=2, column=2, columnspan=2, sticky="w", pady=(10, 0))
        ttk.Checkbutton(frm_cfg, text="处理前创建备份", variable=self.backup_var).grid(row=2, column=4, columnspan=2, sticky="w", pady=(10, 0))
        ttk.Checkbutton(frm_cfg, text="清空目标列旧图片后重做", variable=self.rebuild_var).grid(row=3, column=0, columnspan=2, sticky="w", pady=(10, 0))
        ttk.Checkbutton(frm_cfg, text="忽略缓存，重新下载", variable=self.force_redownload_var).grid(row=3, column=2, columnspan=2, sticky="w", pady=(10, 0))

        frm_ops = ttk.Frame(self.root, padding=(12, 0, 12, 8))
        frm_ops.pack(fill="x")
        ttk.Button(frm_ops, text="开始处理", command=self.start).pack(side="left")
        ttk.Button(frm_ops, text="停止", command=self.stop).pack(side="left", padx=8)
        ttk.Button(frm_ops, text="清空日志", command=self.clear_log).pack(side="left")

        self.progress = ttk.Progressbar(self.root, mode="determinate")
        self.progress.pack(fill="x", padx=12, pady=(0, 8))

        self.status_var = tk.StringVar(value="就绪")
        ttk.Label(self.root, textvariable=self.status_var).pack(fill="x", padx=12)

        self.log_text = tk.Text(self.root, wrap="word", height=28)
        self.log_text.pack(fill="both", expand=True, padx=12, pady=(6, 12))

    def log(self, message: str):
        self.root.after(0, self._append_log, message)

    def _append_log(self, message: str):
        self.log_text.insert("end", message + "\n")
        self.log_text.see("end")

    def clear_log(self):
        self.log_text.delete("1.0", "end")

    def set_status(self, message: str):
        self.root.after(0, lambda m=message: self.status_var.set(m))

    def choose_file(self):
        path = filedialog.askopenfilename(
            title="选择 Excel 文件",
            filetypes=[("Excel 文件", "*.xlsx")]
        )
        if path:
            self.file_var.set(path)
            self.load_sheets()

    def load_sheets(self):
        path = self.file_var.get().strip()
        if not path:
            messagebox.showwarning("提示", "请先选择 Excel 文件。")
            return
        try:
            wb = load_workbook(path, read_only=True)
            sheets = wb.sheetnames
            wb.close()
            self.sheet_combo["values"] = sheets
            if sheets and self.sheet_var.get() not in sheets:
                preferred = next((s for s in sheets if s.startswith("3.16")), sheets[0])
                self.sheet_var.set(preferred)
            self.log(f"已读取工作表：{', '.join(sheets)}")
        except Exception as e:
            messagebox.showerror("错误", f"读取工作表失败：{e}")

    def stop(self):
        self.stop_flag = True
        self.set_status("正在请求停止...")
        self.log("已请求停止，当前行处理完后结束。")

    def start(self):
        if self.worker and self.worker.is_alive():
            messagebox.showwarning("提示", "任务正在运行中。")
            return
        try:
            config = self._collect_config()
        except Exception as e:
            messagebox.showerror("参数错误", str(e))
            return

        self.stop_flag = False
        self.progress["value"] = 0
        self.progress["maximum"] = 1
        self.set_status("准备开始...")
        self.worker = threading.Thread(target=self.run_task, args=(config,), daemon=True)
        self.worker.start()

    def _collect_config(self):
        path = self.file_var.get().strip()
        if not path:
            raise ValueError("请选择 Excel 文件。")
        if not os.path.exists(path):
            raise ValueError("Excel 文件不存在。")
        sheet_name = self.sheet_var.get().strip()
        if not sheet_name:
            raise ValueError("请选择工作表。")
        start_row = int(self.start_row_var.get().strip())
        if start_row < 1:
            raise ValueError("起始行必须大于等于 1。")
        max_width = int(self.max_width_var.get().strip())
        max_height = int(self.max_height_var.get().strip())
        timeout = int(self.timeout_var.get().strip())
        concurrent = int(self.concurrent_var.get().strip())
        if concurrent < 1 or concurrent > 16:
            raise ValueError("并发下载数必须在 1 到 16 之间。")
        return {
            "path": path,
            "sheet_name": sheet_name,
            "url_col": normalize_col(self.url_col_var.get()),
            "target_col": normalize_col(self.target_col_var.get()),
            "start_row": start_row,
            "max_width": max_width,
            "max_height": max_height,
            "timeout": timeout,
            "concurrent": concurrent,
            "copy_mode": bool(self.copy_mode_var.get()),
            "backup": bool(self.backup_var.get()),
            "rebuild": bool(self.rebuild_var.get()),
            "force_redownload": bool(self.force_redownload_var.get()),
        }

    def run_task(self, config: dict):
        err_msg = None
        try:
            path = Path(config["path"])
            output_path = path.with_name(path.stem + "_with_images" + path.suffix) if config["copy_mode"] else path
            backup_path = path.with_name(path.stem + "_backup_before_images" + path.suffix)

            if config["backup"] and not backup_path.exists():
                shutil.copy2(path, backup_path)
                self.log(f"已创建备份：{backup_path}")

            if config["copy_mode"]:
                shutil.copy2(path, output_path)
                self.log(f"将输出到副本：{output_path}")

            wb = load_workbook(output_path)
            if config["sheet_name"] not in wb.sheetnames:
                raise ValueError(f"工作表不存在：{config['sheet_name']}")
            ws = wb[config["sheet_name"]]

            rows = list(range(config["start_row"], ws.max_row + 1))

            if config["rebuild"]:
                removed = self._remove_existing_images_in_column(ws, config["target_col"])
                if removed:
                    self.log(f"已移除目标列旧图片：{removed} 张")
            existing_image_rows = self._get_existing_image_rows(ws, config["target_col"])

            session = requests.Session()
            session.headers.update({"User-Agent": "Mozilla/5.0"})

            ok = 0
            fail = 0
            skip = 0
            tasks = []
            for row in rows:
                url = ws[f"{config['url_col']}{row}"].value
                if not url or not str(url).strip().startswith("http"):
                    skip += 1
                    continue
                url = str(url).strip()
                if not config["rebuild"] and row in existing_image_rows:
                    skip += 1
                    continue
                tasks.append((row, url))

            self.root.after(0, lambda total=max(len(tasks), 1): self.progress.configure(maximum=total, value=0))
            self.log(f"待处理行数：{len(tasks)}，跳过行数：{skip}")

            processed = 0
            with ThreadPoolExecutor(max_workers=config["concurrent"]) as executor:
                future_map = {
                    executor.submit(
                        self._download_and_prepare_image,
                        session,
                        url,
                        config["max_width"],
                        config["max_height"],
                        config["timeout"],
                        config["force_redownload"],
                    ): (row, url)
                    for row, url in tasks
                }
                for future in as_completed(future_map):
                    row, _url = future_map[future]
                    if self.stop_flag:
                        self.log("任务已停止。")
                        break
                    processed += 1
                    try:
                        img_path = future.result()
                        xl_img = XLImage(str(img_path))
                        xl_img.width = config["max_width"]
                        xl_img.height = config["max_height"]
                        ws.add_image(xl_img, f"{config['target_col']}{row}")
                        ws.row_dimensions[row].height = max(72, config["max_height"] * 0.8)
                        ok += 1
                        self._set_progress(processed, f"成功：第 {row} 行")
                    except Exception as e:
                        fail += 1
                        ws[f"{config['target_col']}{row}"] = f"DOWNLOAD_FAIL: {e}"
                        self.log(f"第 {row} 行失败：{e}")
                        self._set_progress(processed, f"失败：第 {row} 行")

            ws.column_dimensions[config["target_col"]].width = max(16, int(config["max_width"] / 6))
            self.log("开始保存输出文件，请等待...")
            self._atomic_save_workbook(wb, output_path)
            self.log(f"处理完成：成功 {ok}，失败 {fail}，跳过 {skip}")
            self.log(f"输出文件：{output_path}")
            self.set_status("处理完成")
            self.root.after(0, lambda p=str(output_path): messagebox.showinfo("完成", f"处理完成。\n文件：{p}"))
            return
        except Exception as e:
            err_msg = str(e)
            self.log(f"运行失败：{e}")
            self.set_status("运行失败")
        if err_msg:
            self.root.after(0, lambda m=err_msg: messagebox.showerror("错误", m))

    def _set_progress(self, value: int, status: str):
        self.root.after(0, lambda v=value: self.progress.configure(value=v))
        self.set_status(status)

    def _remove_existing_images_in_column(self, ws, target_col: str) -> int:
        removed = 0
        target_idx = self._column_index(target_col) - 1
        remaining = []
        for img in getattr(ws, "_images", []):
            anchor = getattr(img, "anchor", None)
            col = None
            if hasattr(anchor, "_from"):
                col = anchor._from.col
            elif hasattr(anchor, "from_"):
                col = anchor.from_.col
            if col == target_idx:
                removed += 1
                continue
            remaining.append(img)
        ws._images = remaining
        return removed

    def _get_existing_image_rows(self, ws, target_col: str) -> set[int]:
        rows = set()
        target_idx = self._column_index(target_col) - 1
        for img in getattr(ws, "_images", []):
            anchor = getattr(img, "anchor", None)
            marker = None
            if hasattr(anchor, "_from"):
                marker = anchor._from
            elif hasattr(anchor, "from_"):
                marker = anchor.from_
            if marker and marker.col == target_idx:
                rows.add(marker.row + 1)
        return rows

    def _download_and_prepare_image(self, session, url: str, max_width: int, max_height: int, timeout: int, force_redownload: bool) -> Path:
        digest = hashlib.md5(url.encode("utf-8")).hexdigest()
        img_path = self.cache_dir / f"{digest}.png"
        if img_path.exists() and not force_redownload:
            return img_path

        resp = session.get(url, timeout=timeout)
        resp.raise_for_status()
        data = io.BytesIO(resp.content)
        with PILImage.open(data) as im:
            im = im.convert("RGB")
            im.thumbnail((max_width, max_height))
            im.save(img_path, format="PNG")
        return img_path

    @staticmethod
    def _atomic_save_workbook(wb, output_path):
        output_path = Path(output_path)
        temp_path = output_path.with_name(f"{output_path.stem}.{uuid.uuid4().hex}.tmp{output_path.suffix}")
        try:
            wb.save(temp_path)
            os.replace(temp_path, output_path)
        finally:
            if temp_path.exists():
                try:
                    temp_path.unlink()
                except Exception:
                    pass

    @staticmethod
    def _column_index(col_name: str) -> int:
        idx = 0
        for ch in col_name:
            idx = idx * 26 + (ord(ch) - ord("A") + 1)
        return idx


if __name__ == "__main__":
    root = tk.Tk()
    app = ExcelImageEmbedApp(root)
    root.mainloop()
