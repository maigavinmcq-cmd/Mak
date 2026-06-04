import json
import os
import re
import subprocess
import sys
import threading
import traceback
import html as html_lib
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import datetime, date, UTC, timezone, timedelta
from urllib.parse import urlparse
from zoneinfo import ZoneInfo

import tkinter as tk
from tkinter import ttk, filedialog, messagebox

import requests
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from playwright.sync_api import sync_playwright


# =========================================================
# 基础配置与通用工具

def parse_date_ymd(s: str) -> date:
    return datetime.strptime(s.strip(), "%Y-%m-%d").date()


def _named_timezone(zone_name: str, fallback_hours: int):
    try:
        return ZoneInfo(zone_name)
    except Exception:
        return timezone(timedelta(hours=fallback_hours))


TIMEZONE_OPTIONS = {
    "UTC+08:00": timezone(timedelta(hours=8)),
    "UTC": UTC,
    "SYSTEM_LOCAL": datetime.now().astimezone().tzinfo or UTC,
    "墨西哥城 (America/Mexico_City)": _named_timezone("America/Mexico_City", -6),
    "坎昆 (America/Cancun)": _named_timezone("America/Cancun", -5),
    "蒂华纳 (America/Tijuana)": _named_timezone("America/Tijuana", -8),
    "奇瓦瓦 (America/Chihuahua)": _named_timezone("America/Chihuahua", -6),
}

TIMEZONE_ALIASES = {
    "新加坡 / 中国 (UTC+08:00)": "UTC+08:00",
    "中国 / 新加坡 (UTC+08:00)": "UTC+08:00",
    "北京时间 (UTC+08:00)": "UTC+08:00",
    "中国标准时间": "UTC+08:00",
    "本机时区": "SYSTEM_LOCAL",
    "系统时区": "SYSTEM_LOCAL",
    "本地时区": "SYSTEM_LOCAL",
    "墨西哥": "墨西哥城 (America/Mexico_City)",
    "墨西哥城": "墨西哥城 (America/Mexico_City)",
    "Mexico City": "墨西哥城 (America/Mexico_City)",
    "America/Mexico_City": "墨西哥城 (America/Mexico_City)",
    "坎昆": "坎昆 (America/Cancun)",
    "Cancun": "坎昆 (America/Cancun)",
    "America/Cancun": "坎昆 (America/Cancun)",
    "蒂华纳": "蒂华纳 (America/Tijuana)",
    "Tijuana": "蒂华纳 (America/Tijuana)",
    "America/Tijuana": "蒂华纳 (America/Tijuana)",
    "奇瓦瓦": "奇瓦瓦 (America/Chihuahua)",
    "Chihuahua": "奇瓦瓦 (America/Chihuahua)",
    "America/Chihuahua": "奇瓦瓦 (America/Chihuahua)",
}


def normalize_timezone_name(name: str) -> str:
    value = (name or "").strip()
    if not value:
        return "UTC+08:00"
    if value in TIMEZONE_OPTIONS:
        return value
    return TIMEZONE_ALIASES.get(value, value)

ACCOUNT_GROUPS_FILE = os.path.abspath("account_groups.json")
UI_SETTINGS_FILE = os.path.abspath("account_video_data_ui.json")
APP_BASE_DIR = os.path.abspath(os.path.dirname(__file__))
CDP_PROFILE_DIR = os.path.join(APP_BASE_DIR, ".cdp_browser_profile")
LOGS_DIR = os.path.join(APP_BASE_DIR, "logs")

# 页面等待参数：在保证稳定性的前提下尽量提升抓取速度。

DETAIL_PAGE_STABILIZE_MS = 250
PROFILE_PAGE_NETWORK_IDLE_TIMEOUT_MS = 1800
PROFILE_SCROLL_PAUSE_MS = 250
PROFILE_SCROLL_NETWORK_IDLE_TIMEOUT_MS = 1000
PROFILE_FINAL_SETTLE_MS = 250
PROFILE_OPEN_SETTLE_MS = 120
API_CAPTURE_MAX_SCROLL_ROUNDS = 20
API_CAPTURE_MIN_SCROLL_ROUNDS = 4
API_CAPTURE_IDLE_LIMIT = 3
ACCOUNT_CRAWL_MAX_WORKERS = 3
DETAIL_FALLBACK_MAX_PER_ACCOUNT = 3


def normalize_username(account: str) -> str:
    """
    支持三种账号输入形式：

    1. username
    2. @username
    3. https://www.tiktok.com/@username
    """
    account = (account or "").strip()
    if not account:
        return ""

    if account.startswith("http://") or account.startswith("https://"):
        parsed = urlparse(account)
        path = parsed.path.strip("/")
        if path.startswith("@"):
            return path[1:].strip()
        raise ValueError(f"无法识别账号链接：{account}")

    return account.lstrip("@").strip()


def check_cdp_endpoint(debug_url: str, timeout: float = 3.0):
    """
    检查远程调试端口是否可连接。
    优先探测 /json/version，以便确认是否为 Chrome/Edge 的 CDP 服务。

    """
    debug_url = (debug_url or "").strip().rstrip("/")
    if not debug_url:
        return False, "调试浏览器地址为空"

    probe_url = f"{debug_url}/json/version"
    try:
        resp = requests.get(probe_url, timeout=timeout)
        if resp.ok:
            return True, None
        return False, f"调试端口已响应，但状态码异常：HTTP {resp.status_code}"
    except requests.exceptions.ConnectionError:
        return False, (
            "无法连接调试浏览器地址。\n"
            "请先使用带远程调试端口的 Chrome/Edge 启动浏览器，例如：\n"
            "chrome.exe --remote-debugging-port=9222\n"
            "或\n"
            "msedge.exe --remote-debugging-port=9222"
        )
    except requests.exceptions.Timeout:
        return False, "连接调试浏览器地址超时，请确认浏览器已启动且端口可访问"
    except Exception as e:
        return False, f"调试浏览器地址检查失败：{e}"
        return False, "连接调试浏览器地址超时，请确认浏览器已启动且端口可访问"

    except Exception as e:
        return False, f"调试浏览器地址检查失败：{e}"


def list_open_tab_urls_via_cdp(debug_url: str):
    out = []
    with sync_playwright() as p:
        browser = p.chromium.connect_over_cdp(debug_url)
        try:
            for context in browser.contexts:
                for page in context.pages:
                    try:
                        out.append({
                            "title": page.title(),
                            "url": page.url,
                        })
                    except Exception:
                        continue
        finally:
            try:
                browser.close()
            except Exception:
                pass
    return out


def find_browser_executable() -> str:
    candidates = [
        os.path.expandvars(r"%ProgramFiles(x86)%\Microsoft\Edge\Application\msedge.exe"),
        os.path.expandvars(r"%ProgramFiles%\Microsoft\Edge\Application\msedge.exe"),
        os.path.expandvars(r"%LocalAppData%\Microsoft\Edge\Application\msedge.exe"),
        os.path.expandvars(r"%ProgramFiles%\Google\Chrome\Application\chrome.exe"),
        os.path.expandvars(r"%ProgramFiles(x86)%\Google\Chrome\Application\chrome.exe"),
        os.path.expandvars(r"%LocalAppData%\Google\Chrome\Application\chrome.exe"),
    ]
    for path in candidates:
        if path and os.path.isfile(path):
            return path
    return ""


def extract_port_from_debug_url(debug_url: str) -> int:
    debug_url = (debug_url or "").strip()
    try:
        parsed = urlparse(debug_url if "://" in debug_url else f"http://{debug_url}")
        return int(parsed.port or 9222)
    except Exception as e:
        raise ValueError(f"无效的调试浏览器地址：{debug_url}") from e


def launch_debug_browser(debug_url: str, log_func=None) -> tuple[bool, str]:
    browser_path = find_browser_executable()
    if not browser_path:
        return False, "未找到 Edge/Chrome 浏览器"

    port = extract_port_from_debug_url(debug_url)
    os.makedirs(CDP_PROFILE_DIR, exist_ok=True)
    cmd = [
        browser_path,
        f"--remote-debugging-port={port}",
        f"--user-data-dir={CDP_PROFILE_DIR}",
        "--no-first-run",
        "--no-default-browser-check",
    ]
    try:
        subprocess.Popen(cmd, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
        if log_func:
            log_func(f"已启动调试浏览器：{os.path.basename(browser_path)} | 端口 {port}")
        return True, browser_path
    except Exception as e:
        return False, f"启动浏览器失败：{e}"


def parse_accounts_from_text(raw_text: str):
    """
    从多行文本中提取账号列表，自动去重并规范化为 username。
    """
    lines = raw_text.replace("\r\n", "\n").replace("\r", "\n").split("\n")
    accounts = []
    seen = set()

    for line in lines:
        line = line.strip()
        if not line:
            continue
        username = normalize_username(line)
        if username and username not in seen:
            seen.add(username)
            accounts.append(username)

    return accounts


def extract_video_id(url: str) -> str:
    if not url:
        return ""
    m = re.search(r"/video/(\d+)", url)
    return m.group(1) if m else ""


def extract_username_from_profile_url(url: str) -> str:
    """
    从 TikTok 主页 URL 中提取用户名，支持带查询参数的链接。
    例如：https://www.tiktok.com/@abc?lang=en
    """
    m = re.search(r"tiktok\.com/@([^/?]+)", url)
    return m.group(1) if m else ""


def ts_to_date(ts_value, target_tz):
    """
    将 Unix 时间戳或 createTime 转换为目标时区下的日期。
    """
    try:
        ts_str = str(ts_value).strip()
        if not ts_str:
            return None
        ts = int(ts_str[:10])
        return datetime.fromtimestamp(ts, UTC).astimezone(target_tz).date()
    except Exception:
        return None


def ts_to_datetime_pair(ts_value, target_tz):
    try:
        ts_str = str(ts_value).strip()
        if not ts_str:
            return None, None
        ts = int(ts_str[:10])
        dt_utc = datetime.fromtimestamp(ts, UTC)
        return dt_utc, dt_utc.astimezone(target_tz)
    except Exception:
        return None, None


def infer_publish_date_from_video_id(video_id: str, target_tz):
    """
    尝试根据 TikTok video_id 的高位时间信息反推出发布日期。
    """
    try:
        vid = int(str(video_id).strip())
        ts = vid >> 32
        return datetime.fromtimestamp(ts, UTC).astimezone(target_tz).date()
    except Exception:
        return None


def infer_video_id_datetime_pair(video_id: str, target_tz):
    try:
        vid = int(str(video_id).strip())
        ts = vid >> 32
        dt_utc = datetime.fromtimestamp(ts, UTC)
        return dt_utc, dt_utc.astimezone(target_tz)
    except Exception:
        return None, None


def visual_len(value: str) -> int:
    total = 0
    for ch in value:
        total += 2 if ord(ch) > 127 else 1
    return total


def auto_fit_columns(ws):
    for column_cells in ws.columns:
        max_length = 0
        col_letter = column_cells[0].column_letter
        for cell in column_cells:
            try:
                value = "" if cell.value is None else str(cell.value)
                width = visual_len(value)
                if width > max_length:
                    max_length = width
            except Exception:
                pass
        ws.column_dimensions[col_letter].width = min(max_length + 2, 60)


def build_account_group_map(account_groups):
    account_to_group = {}
    duplicates = {}

    for group in account_groups or []:
        group_name = (group.get("group_name") or "").strip() or "未命名分组"
        for account in group.get("accounts", []):
            username = normalize_username(account)
            if not username:
                continue
            prev_group = account_to_group.get(username)
            if prev_group and prev_group != group_name:
                duplicates.setdefault(username, {prev_group, group_name}).add(group_name)
                duplicates[username].add(prev_group)
                continue
            account_to_group[username] = group_name

    duplicate_rows = []
    for username, groups in duplicates.items():
        duplicate_rows.append({
            "username": username,
            "groups": sorted(groups),
        })

    duplicate_rows.sort(key=lambda x: x["username"])
    return account_to_group, duplicate_rows


class RunLogWriter:
    def __init__(self, base_dir: str):
        self.base_dir = os.path.abspath(base_dir)
        self.path = ""
        self._lock = threading.Lock()

    def start(self):
        os.makedirs(self.base_dir, exist_ok=True)
        self.path = os.path.join(self.base_dir, f"run_{datetime.now().strftime('%Y%m%d_%H%M%S')}.log")
        with open(self.path, "a", encoding="utf-8") as f:
            f.write("")
        return self.path

    def write(self, msg: str):
        if not self.path:
            return
        line = f"[{datetime.now().strftime('%H:%M:%S')}] {msg}\n"
        with self._lock:
            with open(self.path, "a", encoding="utf-8") as f:
                f.write(line)


def export_to_excel(records, summary_rows, group_summary_rows, output_path, diagnostic_rows=None, log_path: str = ""):
    out_dir = os.path.dirname(os.path.abspath(output_path))
    if out_dir:
        os.makedirs(out_dir, exist_ok=True)
    wb = Workbook()

    ws1 = wb.active
    ws1.title = "视频明细"
    headers1 = [
        "group_name",
        "username",
        "publish_date",
        "video_id",
        "play_count",
        "video_url",
        "author_username",
        "author_match_source",
        "publish_source",
        "publish_confidence",
        "source_tab",
    ]
    ws1.append(headers1)

    for c in ws1[1]:
        c.font = Font(bold=True)
        c.alignment = Alignment(horizontal="center", vertical="center")

    for row in records:
        ws1.append([
            row.get("group_name", ""),
            row.get("username", ""),
            row.get("publish_date", ""),
            row.get("video_id", ""),
            row.get("play_count", ""),
            row.get("video_url", ""),
            row.get("author_username", ""),
            row.get("author_match_source", ""),
            row.get("publish_source", ""),
            row.get("publish_confidence", ""),
            row.get("source_tab", ""),
        ])

    auto_fit_columns(ws1)

    ws2 = wb.create_sheet("账号汇总")
    headers2 = ["group_name", "username", "video_count"]
    ws2.append(headers2)
    for c in ws2[1]:
        c.font = Font(bold=True)
        c.alignment = Alignment(horizontal="center", vertical="center")
    for row in summary_rows:
        ws2.append([
            row.get("group_name", ""),
            row.get("username", ""),
            row.get("video_count", 0),
        ])
    auto_fit_columns(ws2)

    ws3 = wb.create_sheet("分组汇总")
    headers3 = ["group_name", "account_count", "video_count"]
    ws3.append(headers3)
    for c in ws3[1]:
        c.font = Font(bold=True)
        c.alignment = Alignment(horizontal="center", vertical="center")
    for row in group_summary_rows:
        ws3.append([
            row.get("group_name", ""),
            row.get("account_count", 0),
            row.get("video_count", 0),
        ])
    auto_fit_columns(ws3)

    # 新增：置信度汇总
    ws4 = wb.create_sheet("置信度汇总")
    ws4.append(["publish_confidence", "video_count"])
    for c in ws4[1]:
        c.font = Font(bold=True)
        c.alignment = Alignment(horizontal="center", vertical="center")

    confidence_count = {}
    for row in records:
        key = (row.get("publish_confidence") or "").strip() or "UNKNOWN"
        confidence_count[key] = confidence_count.get(key, 0) + 1

    for key in ("HIGH", "MEDIUM", "LOW", "UNKNOWN"):
        ws4.append([key, confidence_count.get(key, 0)])

    auto_fit_columns(ws4)

    ws5 = wb.create_sheet("抓取诊断")
    headers5 = [
        "group_name",
        "username",
        "api_candidate_count",
        "api_payload_hit_count",
        "api_valid_video_count",
        "supplemental_link_count",
        "candidate_to_process_count",
        "author_mismatch_count",
        "duplicate_video_id_count",
        "date_filtered_count",
        "publish_resolved_count",
        "publish_unresolved_count",
        "final_record_count",
        "status",
        "notes",
    ]
    ws5.append(headers5)
    for c in ws5[1]:
        c.font = Font(bold=True)
        c.alignment = Alignment(horizontal="center", vertical="center")
    for row in diagnostic_rows or []:
        ws5.append([
            row.get("group_name", ""),
            row.get("username", ""),
            row.get("api_candidate_count", 0),
            row.get("api_payload_hit_count", 0),
            row.get("api_valid_video_count", 0),
            row.get("supplemental_link_count", 0),
            row.get("candidate_to_process_count", 0),
            row.get("author_mismatch_count", 0),
            row.get("duplicate_video_id_count", 0),
            row.get("date_filtered_count", 0),
            row.get("publish_resolved_count", 0),
            row.get("publish_unresolved_count", 0),
            row.get("final_record_count", 0),
            row.get("status", ""),
            row.get("notes", ""),
        ])
    if log_path:
        ws5.append([])
        ws5.append(["log_path", os.path.abspath(log_path)])
    auto_fit_columns(ws5)
    wb.save(output_path)

# =========================================================
# 页面提取辅助函数


def extract_video_links_from_html(html: str):
    """从主页 HTML 中提取视频详情页链接。"""

    found = set()
    if not html:
        return []

    html = html_lib.unescape(html)
    patterns = [
        r'href="(https://www\.tiktok\.com/@[^"]+/video/\d+)"',
        r"href='(https://www\.tiktok\.com/@[^']+/video/\d+)'",
        r'href="(/@[^"]+/video/\d+)"',
        r"href='(/@[^']+/video/\d+)'",
    ]

    for pattern in patterns:
        for match in re.findall(pattern, html, re.I):
            href = match.strip().split("?")[0]
            if href.startswith("/"):
                href = "https://www.tiktok.com" + href
            if "tiktok.com" in href and "/video/" in href:
                found.add(href)

    return sorted(found)


def extract_video_links_from_dom(page):
    """直接从当前页面 DOM 中提取视频详情页链接。"""

    found = set()
    try:
        hrefs = page.eval_on_selector_all(
            'a[href*="/video/"]',
            """
            els => els
                .map(a => a.href || a.getAttribute('href') || '')
                .filter(Boolean)
            """,
        )
        for href in hrefs:
            href = href.strip().split("?")[0]
            if href.startswith("/"):
                href = "https://www.tiktok.com" + href
            if "tiktok.com" in href and "/video/" in href:
                found.add(href)
    except Exception:
        pass
    return sorted(found)


def extract_video_time_map_from_profile_html(html: str):
    """从主页 HTML/JSON 中提取每个视频的 createTime 映射，返回 video_id -> create_time。"""
    results = {}
    json_candidates = []
    patterns = [
        r'<script id="SIGI_STATE" type="application/json">(.*?)</script>',
        r'<script id="__UNIVERSAL_DATA_FOR_REHYDRATION__" type="application/json">(.*?)</script>',
    ]

    for pattern in patterns:
        json_candidates.extend(re.findall(pattern, html, re.S))

    def walk(obj):
        if isinstance(obj, dict):
            vid = str(obj.get("id", "")).strip() if "id" in obj else ""
            ctime = obj.get("createTime") or obj.get("create_time")
            if vid and ctime and vid not in results:
                results[vid] = str(ctime)
            for value in obj.values():
                walk(value)
        elif isinstance(obj, list):
            for item in obj:
                walk(item)

    for raw_json in json_candidates:
        try:
            walk(json.loads(raw_json))
        except Exception:
            continue

    fallback_patterns = [
        r'"id":"(\d+)".{0,300}?"createTime":"?(\d{10})"?',
        r'"id":"(\d+)".{0,300}?"create_time":"?(\d{10})"?',
        r'"id":(\d+).{0,300}?"createTime":"?(\d{10})"?',
        r'"id":(\d+).{0,300}?"create_time":"?(\d{10})"?',
    ]
    for pattern in fallback_patterns:
        for match in re.findall(pattern, html, re.S):
            if isinstance(match, tuple) and len(match) >= 2:
                vid = str(match[0]).strip()
                ctime = str(match[1]).strip()
                if vid and ctime and vid not in results:
                    results[vid] = ctime

    return results


def _norm_play_count_value(value) -> str:
    if value is None or isinstance(value, bool):
        return ""
    if isinstance(value, (int, float)):
        try:
            return str(int(value))
        except Exception:
            return str(value)
    s = str(value).strip()
    if not s:
        return ""
    m = re.search(r"\d+", s.replace(",", ""))
    return m.group(0) if m else s


def extract_video_play_count_map_from_profile_html(html: str):
    """从主页 HTML/JSON 中提取每个视频的播放量映射，返回 video_id -> play_count。"""
    results = {}
    json_candidates = []
    patterns = [
        r'<script id="SIGI_STATE" type="application/json">(.*?)</script>',
        r'<script id="__UNIVERSAL_DATA_FOR_REHYDRATION__" type="application/json">(.*?)</script>',
    ]

    for pattern in patterns:
        json_candidates.extend(re.findall(pattern, html, re.S))

    def walk(obj):
        if isinstance(obj, dict):
            vid = str(obj.get("id", "")).strip() if "id" in obj else ""
            play_count = ""
            if "playCount" in obj:
                play_count = _norm_play_count_value(obj.get("playCount"))
            elif "play_count" in obj:
                play_count = _norm_play_count_value(obj.get("play_count"))
            elif isinstance(obj.get("stats"), dict):
                stats = obj.get("stats") or {}
                if "playCount" in stats:
                    play_count = _norm_play_count_value(stats.get("playCount"))
                elif "play_count" in stats:
                    play_count = _norm_play_count_value(stats.get("play_count"))
            if vid and play_count and vid not in results:
                results[vid] = play_count
            for value in obj.values():
                walk(value)
        elif isinstance(obj, list):
            for item in obj:
                walk(item)

    for raw_json in json_candidates:
        try:
            walk(json.loads(raw_json))
        except Exception:
            continue

    fallback_patterns = [
        r'"id":"(\d+)".{0,500}?"playCount":"?(\d+)"?',
        r'"id":"(\d+)".{0,500}?"play_count":"?(\d+)"?',
        r'"id":(\d+).{0,500}?"playCount":"?(\d+)"?',
        r'"id":(\d+).{0,500}?"play_count":"?(\d+)"?',
        r'"id":"(\d+)".{0,800}?"stats":\{[^{}]{0,300}?"playCount":"?(\d+)"?',
        r'"id":(\d+).{0,800}?"stats":\{[^{}]{0,300}?"playCount":"?(\d+)"?',
    ]
    for pattern in fallback_patterns:
        for match in re.findall(pattern, html, re.S):
            if isinstance(match, tuple) and len(match) >= 2:
                vid = str(match[0]).strip()
                play_count = str(match[1]).strip()
                if vid and play_count and vid not in results:
                    results[vid] = play_count

    return results


DETAIL_PUBLISH_TIME_KEYS = [
    "publishTime",
    "publish_time",
    "publishedAt",
    "published_at",
    "releaseTime",
    "release_time",
    "datePublished",
    "date_published",
    "publicTime",
    "public_time",
    "goLiveTime",
    "go_live_time",
]


def parse_datetime_candidate(value, target_tz):
    if value is None:
        return None, None

    raw = str(int(value)) if isinstance(value, (int, float)) else str(value).strip()
    if not raw:
        return None, None

    digits = re.sub(r"\D", "", raw)
    try:
        if digits.isdigit() and len(digits) in (10, 13):
            ts = int(digits[:10]) if len(digits) == 10 else int(digits[:13]) / 1000
            dt_utc = datetime.fromtimestamp(ts, UTC)
            return dt_utc, dt_utc.astimezone(target_tz)
    except Exception:
        pass

    normalized = raw.replace("Z", "+00:00")
    for fmt in ("%Y/%m/%d %H:%M", "%Y/%m/%d %H:%M:%S", "%Y-%m-%d %H:%M", "%Y-%m-%d %H:%M:%S"):
        try:
            local_dt = datetime.strptime(raw, fmt).replace(tzinfo=target_tz)
            return local_dt.astimezone(UTC), local_dt
        except Exception:
            pass

    try:
        parsed = datetime.fromisoformat(normalized)
        if parsed.tzinfo is None:
            parsed = parsed.replace(tzinfo=target_tz)
        return parsed.astimezone(UTC), parsed.astimezone(target_tz)
    except Exception:
        return None, None


def extract_json_script_blocks(html: str):
    blocks = []
    patterns = [
        r'<script id="SIGI_STATE" type="application/json">(.*?)</script>',
        r'<script id="__UNIVERSAL_DATA_FOR_REHYDRATION__" type="application/json">(.*?)</script>',
        r'<script[^>]*type="application/json"[^>]*>(.*?)</script>',
    ]
    for pattern in patterns:
        blocks.extend(re.findall(pattern, html, re.S))
    return blocks


def parse_profile_html_bundle(html: str):
    if not html:
        return {
            "links": [],
            "time_map": {},
            "play_count_map": {},
            "script_blocks": [],
        }

    unescaped_html = html_lib.unescape(html)
    script_blocks = extract_json_script_blocks(unescaped_html)
    return {
        "links": extract_video_links_from_html(unescaped_html),
        "time_map": extract_video_time_map_from_profile_html(unescaped_html),
        "play_count_map": extract_video_play_count_map_from_profile_html(unescaped_html),
        "script_blocks": script_blocks,
    }


def extract_candidate_publish_times_from_video_detail_html(html: str, video_id: str):
    """
    高准确率模式：
    - 只接受 JSON 中“明确属于当前 video_id”的时间字段
    - 不再做全页宽松 regex fallback
    """
    candidates = []
    json_candidates = extract_json_script_blocks(html)

    def add_candidate(field_name, raw_value, container_id):
        candidates.append({
            "field": field_name,
            "raw": raw_value,
            "container_id": container_id,
        })

    def walk(obj, matched_video):
        if isinstance(obj, dict):
            current_id = ""
            for key in ("id", "video_id", "itemId", "aweme_id", "group_id"):
                if key in obj and obj.get(key) is not None:
                    current_id = str(obj.get(key)).strip()
                    break

            is_target = matched_video or (current_id == video_id)

            if is_target:
                for field_name in DETAIL_PUBLISH_TIME_KEYS:
                    if field_name in obj and obj.get(field_name) not in (None, ""):
                        add_candidate(field_name, obj.get(field_name), current_id or "detail")
                for fallback_name in ("createTime", "create_time"):
                    if fallback_name in obj and obj.get(fallback_name) not in (None, ""):
                        add_candidate(fallback_name, obj.get(fallback_name), current_id or "detail")

            for value in obj.values():
                walk(value, is_target)

        elif isinstance(obj, list):
            for item in obj:
                walk(item, matched_video)

    for raw_json in json_candidates:
        try:
            walk(json.loads(raw_json), False)
        except Exception:
            continue

    deduped = []
    seen = set()
    for item in candidates:
        key = (item["field"], str(item["raw"]), item["container_id"])
        if key in seen:
            continue
        seen.add(key)
        deduped.append(item)

    return deduped


def choose_best_publish_time_candidate(candidates, target_tz, preferred_container_id=""):
    """
    排序规则：
    1. 明确发布时间字段（HIGH）
    2. createTime（MEDIUM）
    3. 其他弱字段（LOW）
    4. container_id == 当前 video_id 的优先
    """
    best = None
    best_score = -10**9

    for item in candidates:
        utc_dt, local_dt = parse_datetime_candidate(item.get("raw"), target_tz)
        if utc_dt is None or local_dt is None:
            continue

        field_name = str(item.get("field") or "")
        base_field = field_name.split(".")[-1]
        container_id = str(item.get("container_id") or "")

        if base_field in DETAIL_PUBLISH_TIME_KEYS:
            confidence = "HIGH"
            score = 300
        elif base_field in ("createTime", "create_time"):
            confidence = "MEDIUM"
            score = 200
        else:
            confidence = "LOW"
            score = 100

        if preferred_container_id and container_id == preferred_container_id:
            score += 50
        elif container_id and container_id != "detail":
            score += 10

        enriched = {
            "field": field_name,
            "raw": str(item.get("raw")),
            "container_id": container_id,
            "utc_dt": utc_dt,
            "local_dt": local_dt,
            "confidence": confidence,
            "score": score,
        }

        if score > best_score:
            best = enriched
            best_score = score

    return best



def fetch_detail_page_publish_time(context, video_url: str, video_id: str, target_tz, log_func=None):
    detail_page = None
    try:
        detail_page = context.new_page()
        detail_page.goto(video_url, wait_until="domcontentloaded", timeout=15000)
        detail_page.wait_for_timeout(DETAIL_PAGE_STABILIZE_MS)
        html = detail_page.content()

        candidates = extract_candidate_publish_times_from_video_detail_html(html, video_id)
        best = choose_best_publish_time_candidate(
            candidates=candidates,
            target_tz=target_tz,
            preferred_container_id=video_id,
        )

        if best is None:
            if log_func:
                log_func(f"[detail] {video_id} 未命中详情页发布时间候选字段")
            return None

        if log_func:
            log_func(
                f"[detail] {video_id} 命中 {best['field']}={best['raw']} | "
                f"local={best['local_dt'].isoformat()} | confidence={best['confidence']} | "
                f"container={best['container_id']}"
            )

        return {
            "publish_date": best["local_dt"].date(),
            "publish_source": f"detail:{best['field']}",
            "publish_confidence": best["confidence"],
            "debug_parts": [
                f"detail_field={best['field']}",
                f"detail_raw={best['raw']}",
                f"detail_local={best['local_dt'].isoformat()}",
                f"detail_utc={best['utc_dt'].isoformat()}",
                f"detail_confidence={best['confidence']}",
                f"detail_container={best['container_id']}",
            ],
        }
    except Exception as e:
        if log_func:
            log_func(f"[detail] {video_id} 详情页取时间失败：{e}")
        return None
    finally:
        if detail_page is not None:
            try:
                detail_page.close()
            except Exception:
                pass

def looks_like_video_id(value) -> bool:
    s = str(value or "").strip()
    return s.isdigit() and len(s) >= 18


def build_tiktok_video_url(username: str, video_id: str) -> str:
    return f"https://www.tiktok.com/@{username}/video/{video_id}"

# =========================================================
# 高准确率版本：作者校验 + 候选视频构建
# =========================================================

PUBLISH_CONFIDENCE_RANK = {
    "HIGH": 3,
    "MEDIUM": 2,
    "LOW": 1,
}


def normalize_handle_for_match(value: str) -> str:
    """
    用于作者匹配的统一用户名规范化：
    - 去掉 @
    - 统一小写
    - 支持输入 profile url
    """
    if value is None:
        return ""
    raw = str(value).strip()
    if not raw:
        return ""

    if "tiktok.com/@" in raw:
        try:
            return normalize_username(raw).strip().lower()
        except Exception:
            return ""

    return raw.lstrip("@").strip().lower()


def extract_username_from_video_url(video_url: str) -> str:
    """
    从视频链接中提取用户名并规范化
    """
    try:
        return normalize_handle_for_match(extract_username_from_profile_url(video_url))
    except Exception:
        return ""


def extract_canonical_video_url_from_video_obj(video_obj, fallback_username: str, video_id: str) -> str:
    """
    优先使用接口里自带的视频详情链接，拿不到再回退到按 username + video_id 拼接。
    """
    candidates = []

    def push(value):
        if not value or not isinstance(value, str):
            return
        value = value.strip()
        if not value or "/video/" not in value:
            return
        if value.startswith("/"):
            value = "https://www.tiktok.com" + value
        if "tiktok.com" not in value:
            return
        candidates.append(value.split("?")[0])

    def inspect_container(container):
        if not isinstance(container, dict):
            return
        for key in ("shareUrl", "share_url", "url", "itemUrl", "item_url", "videoLink", "video_link"):
            push(container.get(key))

    if isinstance(video_obj, dict):
        inspect_container(video_obj)
        for key in ("author", "authorInfo", "author_info", "video", "itemInfo", "item_info", "shareInfo", "share_info"):
            inspect_container(video_obj.get(key))

    if candidates:
        return candidates[0]
    return build_tiktok_video_url(fallback_username, video_id)


def _push_author_candidate(candidates, raw_value, source):
    username = normalize_handle_for_match(raw_value)
    if username:
        candidates.append({
            "username": username,
            "source": source,
        })


def extract_author_identity_from_video_obj(video_obj):
    """
    从视频对象里尽可能提取作者用户名。
    只接受“像用户名”的字段，不拿 uid / secUid 这类不能直接比对 username 的字段凑数。
    """
    if not isinstance(video_obj, dict):
        return "", ""

    candidates = []

    def inspect_container(container, prefix):
        if isinstance(container, str):
            _push_author_candidate(candidates, container, prefix)
            return
        if not isinstance(container, dict):
            return

        # 最常见的 username 字段
        for key in (
            "uniqueId", "unique_id", "uniqueID",
            "username", "userName", "user_name",
            "authorName", "author_name",
            "handle",
        ):
            if key in container:
                _push_author_candidate(candidates, container.get(key), f"{prefix}.{key}")

        # 有些接口会给 profile/share 链接
        for key in (
            "authorLink", "author_link",
            "profileUrl", "profile_url",
            "shareUrl", "share_url",
            "url",
        ):
            value = container.get(key)
            if value and isinstance(value, str) and "tiktok.com/@" in value:
                _push_author_candidate(candidates, value, f"{prefix}.{key}")

    inspect_container(video_obj, "root")

    for key in (
        "author", "authorInfo", "author_info",
        "user", "owner", "creator",
        "profile", "profileInfo", "profile_info",
    ):
        inspect_container(video_obj.get(key), key)

    # 选优先级最高的候选
    priority = {
        "author.uniqueId": 100,
        "author.unique_id": 100,
        "author.username": 95,
        "author.userName": 95,
        "author.authorName": 90,
        "author_info.uniqueId": 88,
        "authorInfo.uniqueId": 88,
        "user.uniqueId": 85,
        "owner.uniqueId": 85,
        "creator.uniqueId": 85,
        "root.uniqueId": 70,
        "root.username": 65,
        "root.shareUrl": 60,
    }

    best = None
    best_score = -1
    seen = set()

    for item in candidates:
        key = (item["username"], item["source"])
        if key in seen:
            continue
        seen.add(key)
        score = priority.get(item["source"], 10)
        if score > best_score:
            best = item
            best_score = score

    if not best:
        return "", ""

    return best["username"], best["source"]


def build_validated_video_candidate(video_id, video_url, raw_obj, target_username, match_source, play_count=""):
    target_norm = normalize_handle_for_match(target_username)
    author_username = ""
    author_match_source = ""

    if isinstance(raw_obj, dict):
        author_username, author_match_source = extract_author_identity_from_video_obj(raw_obj)

    # API 对象：作者必须明确且匹配
    if raw_obj is not None:
        if not author_username or author_username != target_norm:
            return None
    else:
        # DOM / HTML 补链路：至少 URL 路径里的 @username 要匹配目标账号
        author_username = extract_username_from_video_url(video_url)
        author_match_source = match_source
        if not author_username or author_username != target_norm:
            return None

    return {
        "video_id": str(video_id).strip(),
        "video_url": video_url.strip().split("?")[0],
        "author_username": author_username,
        "author_match_source": author_match_source or match_source,
        "play_count": (play_count or "").strip(),
        "raw_obj": raw_obj,
    }


def build_supplemental_candidates(dom_links, html_links, target_username):
    """
    DOM / HTML 提取到的链接，只保留 URL 上作者 == 目标账号的候选。
    """
    out = {}
    for source_name, links in (("dom_url", dom_links or []), ("html_url", html_links or [])):
        for link in links:
            video_id = extract_video_id(link)
            if not video_id:
                continue
            item = build_validated_video_candidate(
                video_id=video_id,
                video_url=link,
                raw_obj=None,
                target_username=target_username,
                match_source=source_name,
                play_count="",
            )
            if item and item["video_id"] not in out:
                out[item["video_id"]] = item
    return out


def count_candidate_item_list_entries(payload) -> int:
    best = 0

    def walk(obj):
        nonlocal best
        if isinstance(obj, dict):
            for key in ("itemList", "item_list", "aweme_list", "items"):
                value = obj.get(key)
                if isinstance(value, list):
                    best = max(best, len(value))
            for value in obj.values():
                walk(value)
        elif isinstance(obj, list):
            for item in obj:
                walk(item)

    walk(payload)
    return best

def capture_profile_api_responses(page, profile_url: str, log_func=None, refresh_page: bool = False):
    captured = []
    seen_urls = set()

    def on_response(response):
        try:
            url = response.url or ""
            if not url or url in seen_urls:
                return
            ctype = (response.headers or {}).get("content-type", "")
            if "json" not in ctype.lower() and not url.lower().endswith(".json"):
                return
            payload = response.json()
            seen_urls.add(url)
            captured.append({"url": url, "payload": payload})
        except Exception:
            pass

    def candidate_item_list_count():
        best = 0
        for item in captured:
            payload = item.get("payload")
            if payload is not None:
                best = max(best, count_candidate_item_list_entries(payload))
        return best

    page.on("response", on_response)
    try:
        if refresh_page:
            try:
                current_url = (page.url or "").strip()
            except Exception:
                current_url = ""
            try:
                if current_url and current_url == profile_url:
                    page.reload(wait_until="domcontentloaded", timeout=45000)
                else:
                    page.goto(profile_url, wait_until="domcontentloaded", timeout=45000)
            except Exception:
                page.goto(profile_url, wait_until="domcontentloaded", timeout=45000)
            try:
                page.wait_for_load_state("networkidle", timeout=PROFILE_PAGE_NETWORK_IDLE_TIMEOUT_MS)
            except Exception:
                pass
        idle_rounds = 0
        last_item_list_count = -1
        last_scroll_height = -1

        for round_index in range(API_CAPTURE_MAX_SCROLL_ROUNDS):
            try:
                page.mouse.wheel(0, 1800)
                page.wait_for_timeout(PROFILE_SCROLL_PAUSE_MS)
            except Exception:
                break

            current_item_list_count = candidate_item_list_count()
            try:
                current_scroll_height = int(page.evaluate("() => document.documentElement.scrollHeight || document.body.scrollHeight || 0"))
            except Exception:
                current_scroll_height = last_scroll_height

            item_list_stable = current_item_list_count == last_item_list_count
            scroll_height_stable = current_scroll_height == last_scroll_height
            if item_list_stable and scroll_height_stable:
                idle_rounds += 1
            else:
                idle_rounds = 0
                last_item_list_count = current_item_list_count
                last_scroll_height = current_scroll_height

            try:
                page.wait_for_load_state("networkidle", timeout=PROFILE_SCROLL_NETWORK_IDLE_TIMEOUT_MS)
            except Exception:
                pass

            if log_func:
                log_func(
                    f"[api] 滚动第 {round_index + 1} 轮 | JSON={len(captured)} | "
                    f"item_list={candidate_item_list_count()} | scroll_height={current_scroll_height} | idle={idle_rounds}"
                )
            if round_index + 1 >= API_CAPTURE_MIN_SCROLL_ROUNDS and idle_rounds >= API_CAPTURE_IDLE_LIMIT:
                break

        try:
            page.wait_for_timeout(PROFILE_FINAL_SETTLE_MS)
        except Exception:
            pass

        if log_func:
            log_func(f"[api] 捕获到 JSON 响应 {len(captured)} 个")
        return captured
    finally:
        try:
            page.remove_listener("response", on_response)
        except Exception:
            pass


def _iter_video_like_objects(obj):
    if isinstance(obj, dict):
        video_id = ""
        for key in ("id", "aweme_id", "video_id", "itemId", "group_id"):
            value = obj.get(key)
            if looks_like_video_id(value):
                video_id = str(value).strip()
                break

        has_time = any(obj.get(k) not in (None, "") for k in DETAIL_PUBLISH_TIME_KEYS + ["createTime", "create_time"])
        has_stats = "stats" in obj or "playCount" in obj or "play_count" in obj
        has_author_hint = any(
            obj.get(k) not in (None, "")
            for k in ("author", "authorInfo", "author_info", "user", "owner", "creator", "shareUrl", "share_url", "url")
        )
        if video_id and (has_time or has_stats or has_author_hint):
            yield obj

        for value in obj.values():
            yield from _iter_video_like_objects(value)
    elif isinstance(obj, list):
        for item in obj:
            yield from _iter_video_like_objects(item)


def find_candidate_video_list_payloads(responses, username: str):
    candidates = []
    username = (username or "").strip().lower()

    for item in responses:
        payload = item.get("payload")
        if payload is None:
            continue

        score = 0
        url = (item.get("url") or "").lower()
        if username and username in url:
            score += 2
        if any(token in url for token in ("post", "item", "aweme", "video", "profile")):
            score += 2

        video_count = 0
        for _ in _iter_video_like_objects(payload):
            video_count += 1
            if video_count >= 5:
                break

        if video_count:
            score += min(video_count, 10)

        if video_count:
            candidates.append({
                "url": item.get("url", ""),
                "payload": payload,
                "score": score,
                "video_count": video_count,
            })

    candidates.sort(key=lambda x: (x["score"], x.get("video_count", 0)), reverse=True)
    return candidates


def extract_cursor_from_item_list_url(url: str) -> int:
    try:
        parsed = urlparse(url)
        m = re.search(r"(?:^|&)cursor=(\d+)(?:&|$)", parsed.query)
        if m:
            return int(m.group(1))
    except Exception:
        pass
    return 0


def extract_play_count_from_video_obj(obj) -> str:
    if not isinstance(obj, dict):
        return ""

    value = obj.get("playCount")
    if value in (None, ""):
        value = obj.get("play_count")
    stats = obj.get("stats") if isinstance(obj.get("stats"), dict) else {}
    if value in (None, ""):
        value = stats.get("playCount")
    if value in (None, ""):
        value = stats.get("play_count")
    return _norm_play_count_value(value)


def extract_videos_from_payload(payload, target_username: str):
    """
    只保留“作者明确 == 目标账号”的 API 视频对象。
    高准确率优先：作者不明确，宁可丢，也不误收。
    """
    extracted = []
    seen = set()
    diagnostics = {
        "video_like_seen": 0,
        "duplicate_video_id": 0,
        "missing_video_id": 0,
        "author_mismatch": 0,
        "accepted": 0,
    }

    for obj in _iter_video_like_objects(payload):
        diagnostics["video_like_seen"] += 1
        video_id = ""
        for key in ("id", "aweme_id", "video_id", "itemId", "group_id"):
            value = obj.get(key)
            if looks_like_video_id(value):
                video_id = str(value).strip()
                break

        if not video_id:
            diagnostics["missing_video_id"] += 1
            continue

        if video_id in seen:
            diagnostics["duplicate_video_id"] += 1
            continue

        play_count = extract_play_count_from_video_obj(obj)
        candidate = build_validated_video_candidate(
            video_id=video_id,
            video_url=extract_canonical_video_url_from_video_obj(obj, target_username, video_id),
            raw_obj=obj,
            target_username=target_username,
            match_source="api_author",
            play_count=play_count,
        )
        if candidate is None:
            diagnostics["author_mismatch"] += 1
            continue

        seen.add(video_id)
        extracted.append(candidate)
        diagnostics["accepted"] += 1

    return extracted, diagnostics


def resolve_publish_time_from_video_obj(video_obj, target_tz):
    raw_obj = (video_obj or {}).get("raw_obj") or {}
    candidates = []

    for field_name in DETAIL_PUBLISH_TIME_KEYS + ["createTime", "create_time"]:
        if field_name in raw_obj and raw_obj.get(field_name) not in (None, ""):
            candidates.append({
                "field": field_name,
                "raw": raw_obj.get(field_name),
                "container_id": (video_obj or {}).get("video_id", ""),
            })

    stats = raw_obj.get("stats")
    if isinstance(stats, dict):
        for field_name in DETAIL_PUBLISH_TIME_KEYS + ["createTime", "create_time"]:
            if field_name in stats and stats.get(field_name) not in (None, ""):
                candidates.append({
                    "field": f"stats.{field_name}",
                    "raw": stats.get(field_name),
                    "container_id": (video_obj or {}).get("video_id", ""),
                })

    best = choose_best_publish_time_candidate(
        candidates=candidates,
        target_tz=target_tz,
        preferred_container_id=(video_obj or {}).get("video_id", ""),
    )
    if best is None:
        return None

    return {
        "publish_date": best["local_dt"].date(),
        "publish_source": f"api:{best['field']}",
        "publish_confidence": best["confidence"],
        "debug_parts": [
            f"api_field={best['field']}",
            f"api_raw={best['raw']}",
            f"api_local={best['local_dt'].isoformat()}",
            f"api_utc={best['utc_dt'].isoformat()}",
            f"api_confidence={best['confidence']}",
        ],
    }

# =========================================================
# 抓取主流程
class TikTokOpenTabsCrawler:
    def __init__(
        self,
        log_func=None,
        stop_flag_func=None,
        refresh_existing: bool = True,
        allow_detail_fallback: bool = False,
        max_workers: int = ACCOUNT_CRAWL_MAX_WORKERS,
        progress_callback=None,
    ):
        self.log_func = log_func or (lambda msg: None)
        self.stop_flag_func = stop_flag_func or (lambda: False)
        self.refresh_existing = bool(refresh_existing)
        self.allow_detail_fallback = bool(allow_detail_fallback)
        self.max_workers = max(1, int(max_workers or 1))
        self.progress_callback = progress_callback or (lambda current, total, username="", phase="": None)

    def log(self, msg: str):
        self.log_func(msg)

    def report_progress(self, current: int, total: int, username: str = "", phase: str = ""):
        try:
            self.progress_callback(current, total, username, phase)
        except Exception:
            pass

    def log_stage(self, stage: str, message: str):
        self.log(f"[{stage}] {message}")

    def should_stop(self):
        return self.stop_flag_func()

    @staticmethod
    def empty_result():
        return [], [], []

    @staticmethod
    def build_account_diagnostic(username: str, group_name: str = ""):
        return {
            "group_name": group_name or "未分组",
            "username": username,
            "api_candidate_count": 0,
            "api_payload_hit_count": 0,
            "api_valid_video_count": 0,
            "supplemental_link_count": 0,
            "candidate_to_process_count": 0,
            "author_mismatch_count": 0,
            "duplicate_video_id_count": 0,
            "date_filtered_count": 0,
            "publish_resolved_count": 0,
            "publish_unresolved_count": 0,
            "final_record_count": 0,
            "status": "",
            "notes": "",
        }

    def resolve_publish_date(
            self,
            context,
            video_url: str,
            video_id: str,
            api_video,
            time_map,
            target_tz,
            allow_detail_fallback: bool | None = None,
    ):
        debug_parts = [f"video_id={video_id}"]
        allow_detail = self.allow_detail_fallback if allow_detail_fallback is None else bool(allow_detail_fallback)

        api_publish = resolve_publish_time_from_video_obj(api_video or {}, target_tz)
        detail_publish = None

        # 先返回 API 里的 HIGH
        if api_publish is not None:
            debug_parts.extend(api_publish["debug_parts"])
            if api_publish["publish_confidence"] == "HIGH":
                return {
                    "publish_date": api_publish["publish_date"],
                    "publish_source": api_publish["publish_source"],
                    "publish_confidence": api_publish["publish_confidence"],
                    "debug_parts": debug_parts,
                }

        # API 里只有 MEDIUM 时，也优先于 HTML createTime
        if api_publish is not None:
            return {
                "publish_date": api_publish["publish_date"],
                "publish_source": api_publish["publish_source"],
                "publish_confidence": api_publish["publish_confidence"],
                "debug_parts": debug_parts,
            }

        # HTML createTime
        if video_id in time_map:
            raw_ts = time_map[video_id]
            utc_dt, local_dt = ts_to_datetime_pair(raw_ts, target_tz)
            debug_parts.append(f"html_createTime={raw_ts}")
            if utc_dt is not None:
                debug_parts.append(f"html_createTime_utc={utc_dt.isoformat()}")
                debug_parts.append(f"html_createTime_local={local_dt.isoformat()}")
                debug_parts.append(f"html_createTime_local_date={local_dt.date().isoformat()}")

            publish_date = ts_to_date(raw_ts, target_tz)
            if publish_date is not None:
                return {
                    "publish_date": publish_date,
                    "publish_source": "html:createTime",
                    "publish_confidence": "MEDIUM",
                    "debug_parts": debug_parts,
                }

        # detail 只对 API/HTML 都没拿到时间的视频触发
        if allow_detail and context is not None and video_url:
            detail_publish = fetch_detail_page_publish_time(
                context=context,
                video_url=video_url,
                video_id=video_id,
                target_tz=target_tz,
                log_func=self.log,
            )
            if detail_publish is not None:
                detail_debug_parts = list(detail_publish["debug_parts"])
                # detail HIGH 直接返回
                if detail_publish["publish_confidence"] == "HIGH":
                    return {
                        "publish_date": detail_publish["publish_date"],
                        "publish_source": detail_publish["publish_source"],
                        "publish_confidence": detail_publish["publish_confidence"],
                        "debug_parts": debug_parts + detail_debug_parts,
                    }
        else:
            detail_debug_parts = []

        # detail 只有 MEDIUM 时，在这里返回
        if detail_publish is not None:
            return {
                "publish_date": detail_publish["publish_date"],
                "publish_source": detail_publish["publish_source"],
                "publish_confidence": detail_publish["publish_confidence"],
                "debug_parts": debug_parts + list(detail_publish["debug_parts"]),
            }

        # 最后才用 video_id 推断
        utc_dt, local_dt = infer_video_id_datetime_pair(video_id, target_tz)
        if utc_dt is not None:
            debug_parts.append(f"video_id_utc={utc_dt.isoformat()}")
            debug_parts.append(f"video_id_local={local_dt.isoformat()}")
            debug_parts.append(f"video_id_local_date={local_dt.date().isoformat()}")

        publish_date = infer_publish_date_from_video_id(video_id, target_tz)
        if publish_date is not None:
            return {
                "publish_date": publish_date,
                "publish_source": "video_id",
                "publish_confidence": "LOW",
                "debug_parts": debug_parts,
            }

        return {
            "publish_date": None,
            "publish_source": "",
            "publish_confidence": "",
            "debug_parts": debug_parts,
        }
    def deduplicate_records(self, records):
        """
        优先按 video_id 去重，取不到再按 video_url 去重
        """
        unique = []
        seen = set()

        for row in records:
            video_id = (row.get("video_id") or "").strip()
            video_url = (row.get("video_url") or "").strip()
            key = f"id:{video_id}" if video_id else f"url:{video_url}"

            if key in seen:
                continue

            seen.add(key)
            unique.append(row)

        return unique

    def build_summary(self, records, account_groups, account_order):
        count_map = {}
        for row in records:
            username = row.get("username", "")
            count_map[username] = count_map.get(username, 0) + 1

        summary = []
        if account_groups:
            for group in account_groups:
                group_name = (group.get("group_name") or "").strip() or "未命名分组"
                for acc in group.get("accounts", []):
                    summary.append({
                        "group_name": group_name,
                        "username": acc,
                        "video_count": count_map.get(acc, 0),
                    })
        else:
            ordered_accounts = account_order or sorted(count_map.keys())
            for acc in ordered_accounts:
                summary.append({
                    "group_name": "未分组",
                    "username": acc,
                    "video_count": count_map.get(acc, 0),
                })

        return summary

    def build_group_summary(self, records, account_groups):
        count_map = {}
        for row in records:
            group_name = (row.get("group_name") or "").strip() or "未分组"
            count_map[group_name] = count_map.get(group_name, 0) + 1

        summary = []
        if account_groups:
            for group in account_groups:
                group_name = (group.get("group_name") or "").strip() or "未命名分组"
                summary.append({
                    "group_name": group_name,
                    "account_count": len(group.get("accounts", [])),
                    "video_count": count_map.get(group_name, 0),
                })
        else:
            summary.append({
                "group_name": "未分组",
                "account_count": len(sorted({row.get('username', '') for row in records if row.get('username', '')})),
                "video_count": len(records),
            })

        return summary

    def ensure_target_account_pages_open(self, browser, target_accounts, refresh_existing: bool = True):
        normalized_accounts = []
        for item in (target_accounts or []):
            username = normalize_username(item)
            if username and username not in normalized_accounts:
                normalized_accounts.append(username)
        if not normalized_accounts:
            return 0

        opened = set()
        opened_pages = {}
        first_context = browser.contexts[0] if browser.contexts else browser.new_context()

        for context in browser.contexts:
            for page in context.pages:
                try:
                    username = extract_username_from_profile_url(page.url)
                except Exception:
                    username = ""
                if username:
                    opened.add(username)
                    opened_pages.setdefault(username, page)

        refresh_targets = [u for u in normalized_accounts if u in opened_pages]
        if refresh_existing and refresh_targets:
            self.log(f"检测到 {len(refresh_targets)} 个已打开账号主页，开始逐个刷新。")
            for idx, username in enumerate(refresh_targets, start=1):
                if self.should_stop():
                    self.log("检测到停止信号，终止自动刷新主页。")
                    break
                page = opened_pages.get(username)
                if page is None:
                    continue
                self.log(f"[刷新主页 {idx}/{len(refresh_targets)}] https://www.tiktok.com/@{username}")
                try:
                    page.reload(wait_until="domcontentloaded", timeout=45000)
                    try:
                        page.wait_for_load_state("networkidle", timeout=PROFILE_PAGE_NETWORK_IDLE_TIMEOUT_MS)
                    except Exception:
                        pass
                    try:
                        page.wait_for_timeout(PROFILE_OPEN_SETTLE_MS)
                    except Exception:
                        pass
                except Exception as e:
                    self.log(f"[刷新失败] @{username} - {e}")
        elif refresh_targets:
            self.log(f"检测到 {len(refresh_targets)} 个已打开账号主页，按设置跳过刷新。")

        missing = [u for u in normalized_accounts if u not in opened]
        if not missing:
            self.log("待抓取账号主页已全部就绪，无需补开。")
            return 0

        self.log(f"检测到 {len(missing)} 个待抓取账号主页未打开，开始自动补开。")
        opened_count = 0
        for idx, username in enumerate(missing, start=1):
            if self.should_stop():
                self.log("检测到停止信号，终止自动刷新主页。")
                break

            url = f"https://www.tiktok.com/@{username}"
            self.log(f"[补开主页 {idx}/{len(missing)}] {url}")
            page = first_context.new_page()
            try:
                page.goto(url, wait_until="domcontentloaded", timeout=45000)
                try:
                    page.wait_for_load_state("networkidle", timeout=PROFILE_PAGE_NETWORK_IDLE_TIMEOUT_MS)
                except Exception:
                    pass
                try:
                    page.wait_for_timeout(PROFILE_OPEN_SETTLE_MS)
                except Exception:
                    pass
                opened_count += 1
            except Exception as e:
                self.log(f"[刷新失败] @{username} - {e}")
                try:
                    page.close()
                except Exception:
                    pass

        self.log(f"自动补开主页完成：成功 {opened_count} 个。")
        return opened_count

    def collect_profile_targets(self, browser, target_accounts=None):
        target_set = set(target_accounts or [])
        page_count = 0
        profile_targets = []
        seen_usernames = set()

        for context in browser.contexts:
            for page in context.pages:
                page_count += 1
                try:
                    url = page.url
                except Exception:
                    continue

                if "tiktok.com/@" not in url or "/video/" in url:
                    continue

                username = extract_username_from_profile_url(url)
                if not username or username in seen_usernames:
                    continue

                if target_set and username not in target_set:
                    continue

                seen_usernames.add(username)
                profile_targets.append({
                    "username": username,
                    "url": url,
                })

        return profile_targets, page_count

    @staticmethod
    def find_profile_page(browser, username: str, preferred_url: str = ""):
        for context in browser.contexts:
            for page in context.pages:
                try:
                    page_url = page.url
                except Exception:
                    continue
                if preferred_url and page_url == preferred_url:
                    return context, page
                if extract_username_from_profile_url(page_url) == username:
                    return context, page

        first_context = browser.contexts[0] if browser.contexts else browser.new_context()
        return first_context, None

    def collect_single_profile_api_first(
            self,
            debug_url: str,
            username: str,
            url: str,
            start_d: date,
            end_d: date,
            target_tz,
            account_to_group=None,
    ):
        account_to_group = account_to_group or {}
        target_username_norm = normalize_handle_for_match(username)
        group_name = account_to_group.get(username, "未分组")

        with sync_playwright() as p:
            browser = p.chromium.connect_over_cdp(debug_url)
            try:
                context, page = self.find_profile_page(browser, username, preferred_url=url)
                if page is None:
                    page = context.new_page()
                    page.goto(url, wait_until="domcontentloaded", timeout=45000)
                    try:
                        page.wait_for_load_state("networkidle", timeout=PROFILE_PAGE_NETWORK_IDLE_TIMEOUT_MS)
                    except Exception:
                        pass
                    try:
                        page.wait_for_timeout(PROFILE_OPEN_SETTLE_MS)
                    except Exception:
                        pass

                self.log("=" * 72)
                self.log(f"处理已打开标签页：{username}")
                self.log(f"标签页 URL：{url}")

                try:
                    title = page.title()
                    self.log(f"页面标题：{title}")
                except Exception:
                    pass

                def run_single_attempt(refresh_page: bool, attempt_label: str):
                    records = []
                    diagnostic = self.build_account_diagnostic(username, group_name)

                    if refresh_page:
                        self.log(f"[{username}] 首次 no_candidates，刷新主页后重试一次")

                    profile_api_responses = capture_profile_api_responses(
                        page,
                        url,
                        log_func=self.log,
                        refresh_page=refresh_page,
                    )
                    api_candidates = find_candidate_video_list_payloads(profile_api_responses, username)
                    api_candidates.sort(key=lambda x: (extract_cursor_from_item_list_url(x["url"]), -x["score"]))
                    diagnostic["api_candidate_count"] = len(profile_api_responses)
                    diagnostic["api_payload_hit_count"] = len(api_candidates)
                    self.log(
                        f"[{username}] API 候选响应数：{len(profile_api_responses)} | 命中 payload 数：{len(api_candidates)}")

                    api_video_map = {}
                    for candidate in api_candidates:
                        extracted, api_diag = extract_videos_from_payload(candidate["payload"], username)
                        diagnostic["author_mismatch_count"] += int(api_diag.get("author_mismatch", 0) or 0)
                        diagnostic["duplicate_video_id_count"] += int(api_diag.get("duplicate_video_id", 0) or 0)
                        if not extracted:
                            self.log(
                                f"[{username}] API 候选未产出有效视频 | score={candidate['score']} | "
                                f"video_like={api_diag.get('video_like_seen', 0)} | "
                                f"dup={api_diag.get('duplicate_video_id', 0)} | "
                                f"author_mismatch={api_diag.get('author_mismatch', 0)} | "
                                f"url={candidate['url']}"
                            )
                            continue
                        self.log(
                            f"[{username}] API 候选命中 | score={candidate['score']} | "
                            f"video_like={api_diag.get('video_like_seen', candidate.get('video_count', 0))} | "
                            f"valid_videos={len(extracted)} | "
                            f"dup={api_diag.get('duplicate_video_id', 0)} | "
                            f"author_mismatch={api_diag.get('author_mismatch', 0)} | "
                            f"url={candidate['url']}"
                        )
                        for item in extracted:
                            if item["video_id"] not in api_video_map:
                                api_video_map[item["video_id"]] = item
                    diagnostic["api_valid_video_count"] = len(api_video_map)

                    try:
                        html = page.content()
                    except Exception as e:
                        self.log(f"[{username}] 读取页面 HTML 失败：{e}")
                        html = ""

                    dom_links = extract_video_links_from_dom(page)
                    html_bundle = parse_profile_html_bundle(html)
                    html_links = html_bundle["links"]
                    time_map = html_bundle["time_map"]
                    play_count_map = html_bundle["play_count_map"]

                    supplemental_map = build_supplemental_candidates(dom_links, html_links, username)
                    diagnostic["supplemental_link_count"] = len(supplemental_map)

                    candidates_to_process = []
                    for item in api_video_map.values():
                        candidates_to_process.append(item)
                    for video_id, item in supplemental_map.items():
                        if video_id not in api_video_map:
                            candidates_to_process.append(item)
                    diagnostic["candidate_to_process_count"] = len(candidates_to_process)

                    self.log(f"[{username}] API 提取并通过作者校验的视频数：{len(api_video_map)}")
                    self.log(f"[{username}] DOM 提取视频链接数：{len(dom_links)}")
                    self.log(f"[{username}] HTML 提取视频链接数：{len(html_links)}")
                    self.log(f"[{username}] DOM/HTML 通过 URL 作者校验的视频数：{len(supplemental_map)}")
                    self.log(f"[{username}] 最终待处理视频数：{len(candidates_to_process)}")

                    matched_count = 0
                    confidence_counter = {"HIGH": 0, "MEDIUM": 0, "LOW": 0}
                    detail_fallback_used = 0

                    for item in candidates_to_process:
                        if self.should_stop():
                            self.log("检测到停止信号，终止任务。")
                            break

                        video_id = item["video_id"]
                        link = item["video_url"]
                        api_video = api_video_map.get(video_id)
                        should_allow_detail_for_video = (
                            self.allow_detail_fallback
                            and detail_fallback_used < DETAIL_FALLBACK_MAX_PER_ACCOUNT
                            and api_video is None
                            and video_id not in time_map
                        )

                        publish_meta = self.resolve_publish_date(
                            context=context,
                            video_url=link,
                            video_id=video_id,
                            api_video=api_video,
                            time_map=time_map,
                            target_tz=target_tz,
                            allow_detail_fallback=should_allow_detail_for_video,
                        )
                        publish_date = publish_meta["publish_date"]
                        publish_source = publish_meta["publish_source"]
                        publish_confidence = publish_meta["publish_confidence"]
                        debug_parts = list(publish_meta["debug_parts"])
                        if publish_source.startswith("detail:"):
                            detail_fallback_used += 1

                        if publish_date is None:
                            diagnostic["publish_unresolved_count"] += 1
                            continue
                        diagnostic["publish_resolved_count"] += 1

                        play_count = (
                                ((api_video_map.get(video_id) or {}).get("play_count") or "").strip()
                                or (play_count_map.get(video_id) or "").strip()
                        )

                        debug_parts.append(f"author_username={item.get('author_username', '')}")
                        debug_parts.append(f"author_match_source={item.get('author_match_source', '')}")
                        debug_parts.append(f"publish_source={publish_source or 'unknown'}")
                        debug_parts.append(f"publish_confidence={publish_confidence or 'unknown'}")
                        debug_parts.append(f"publish_date={publish_date.isoformat()}")
                        if play_count:
                            debug_parts.append(f"play_count={play_count}")

                        self.log(f"[{username}] 日期诊断 | " + " | ".join(debug_parts))

                        if publish_date < start_d or publish_date > end_d:
                            diagnostic["date_filtered_count"] += 1
                            continue

                        if normalize_handle_for_match(item.get("author_username")) != target_username_norm:
                            self.log(
                                f"[{username}] 跳过作者不匹配视频：video_id={video_id} author={item.get('author_username', '')}")
                            diagnostic["author_mismatch_count"] += 1
                            continue

                        records.append({
                            "group_name": group_name,
                            "username": username,
                            "publish_date": publish_date.isoformat(),
                            "video_id": video_id,
                            "play_count": play_count,
                            "video_url": link,
                            "author_username": item.get("author_username", ""),
                            "author_match_source": item.get("author_match_source", ""),
                            "publish_source": publish_source,
                            "publish_confidence": publish_confidence,
                            "source_tab": url,
                        })
                        matched_count += 1
                        confidence_counter[publish_confidence] = confidence_counter.get(publish_confidence, 0) + 1

                    diagnostic["final_record_count"] = matched_count
                    if matched_count > 0:
                        diagnostic["status"] = "ok"
                    elif diagnostic["candidate_to_process_count"] == 0:
                        diagnostic["status"] = "no_candidates"
                    elif diagnostic["publish_resolved_count"] == 0:
                        diagnostic["status"] = "publish_unresolved"
                    elif diagnostic["date_filtered_count"] >= diagnostic["publish_resolved_count"]:
                        diagnostic["status"] = "date_filtered"
                    else:
                        diagnostic["status"] = "zero_records"
                    diagnostic["notes"] = (
                        f"attempt={attempt_label} | "
                        f"api_valid={diagnostic['api_valid_video_count']} | "
                        f"supplemental={diagnostic['supplemental_link_count']} | "
                        f"resolved={diagnostic['publish_resolved_count']} | "
                        f"detail_used={detail_fallback_used}"
                    )
                    self.log(f"[{username}] 日期范围内命中数量：{matched_count}")
                    self.log(
                        f"[{username}] 置信度分层 | "
                        f"HIGH={confidence_counter.get('HIGH', 0)} | "
                        f"MEDIUM={confidence_counter.get('MEDIUM', 0)} | "
                        f"LOW={confidence_counter.get('LOW', 0)}"
                    )
                    self.log(
                        f"[{username}] 账号诊断汇总 | status={diagnostic['status']} | "
                        f"api_candidate={diagnostic['api_candidate_count']} | "
                        f"payload_hit={diagnostic['api_payload_hit_count']} | "
                        f"api_valid={diagnostic['api_valid_video_count']} | "
                        f"supplemental={diagnostic['supplemental_link_count']} | "
                        f"to_process={diagnostic['candidate_to_process_count']} | "
                        f"author_mismatch={diagnostic['author_mismatch_count']} | "
                        f"date_filtered={diagnostic['date_filtered_count']} | "
                        f"final={diagnostic['final_record_count']}"
                    )
                    return records, diagnostic

                records, diagnostic = run_single_attempt(refresh_page=False, attempt_label="initial")
                if diagnostic.get("status") == "no_candidates":
                    records_retry, diagnostic_retry = run_single_attempt(refresh_page=True, attempt_label="retry_after_refresh")
                    if diagnostic_retry.get("status") == "no_candidates":
                        self.log(f"[{username}] 刷新后仍然 no_candidates，跳过当前账号")
                    return records_retry, diagnostic_retry
                return records, diagnostic
            finally:
                try:
                    browser.close()
                except Exception:
                    pass

    def collect_from_open_tabs(
        self,
        debug_url: str,
        start_d: date,
        end_d: date,
        target_tz,
        target_accounts=None,
        account_to_group=None,
        account_groups=None,
    ):
        target_accounts = target_accounts or []
        account_to_group = account_to_group or {}
        account_groups = account_groups or []
        target_set = set(target_accounts)

        all_records = []
        discovered_accounts = []
        skipped_tabs = []

        with sync_playwright() as p:
            self.log(f"连接浏览器：{debug_url}")
            browser = p.chromium.connect_over_cdp(debug_url)

            try:
                self.ensure_target_account_pages_open(browser, target_accounts, refresh_existing=self.refresh_existing)
                page_count = 0

                for context in browser.contexts:
                    for page in context.pages:
                        if self.should_stop():
                            self.log("检测到停止信号，终止任务")
                            return self.empty_result()

                        page_count += 1
                        try:
                            url = page.url
                        except Exception:
                            continue

                        if "tiktok.com/@" not in url or "/video/" in url:
                            skipped_tabs.append(url)
                            continue

                        username = extract_username_from_profile_url(url)
                        if not username:
                            skipped_tabs.append(url)
                            continue

                        if target_set and username not in target_set:
                            continue

                        discovered_accounts.append(username)
                        self.log("=" * 72)
                        self.log(f"处理已打开标签页：{username}")
                        self.log(f"标签页 URL：{url}")

                        try:
                            title = page.title()
                            self.log(f"页面标题：{title}")
                        except Exception:
                            pass

                        try:
                            html = page.content()
                        except Exception as e:
                            self.log(f"[{username}] 读取页面 HTML 失败：{e}")
                            continue

                        dom_links = extract_video_links_from_dom(page)
                        html_bundle = parse_profile_html_bundle(html)
                        html_links = html_bundle["links"]
                        time_map = html_bundle["time_map"]
                        play_count_map = html_bundle["play_count_map"]
                        all_links = sorted(set(dom_links) | set(html_links))

                        self.log(f"[{username}] DOM 提取视频链接数：{len(dom_links)}")
                        self.log(f"[{username}] HTML 提取视频链接数：{len(html_links)}")
                        self.log(f"[{username}] 合并后视频链接数：{len(all_links)}")
                        self.log(f"[{username}] HTML 中显式带时间戳的视频数：{len(time_map)}")
                        self.log(f"[{username}] HTML 中提取到播放量的视频数：{len(play_count_map)}")

                        matched_count = 0
                        has_explicit_time_count = 0
                        has_inferred_time_count = 0
                        has_detail_time_count = 0

                        for link in all_links:
                            if self.should_stop():
                                self.log("检测到停止信号，终止任务")
                                return self.empty_result()

                            video_id = extract_video_id(link)
                            if not video_id:
                                continue

                            publish_date, publish_source, debug_parts = self.resolve_publish_date(
                                context=context,
                                video_url=link,
                                video_id=video_id,
                                api_video=None,
                                time_map=time_map,
                                target_tz=target_tz,
                            )
                            if publish_date is None:
                                continue

                            if publish_source == "createTime":
                                has_explicit_time_count += 1
                            elif publish_source == "video_id":
                                has_inferred_time_count += 1
                            elif publish_source.startswith("detail:"):
                                has_detail_time_count += 1

                            play_count = (play_count_map.get(video_id) or "").strip()
                            debug_parts.append(f"publish_source={publish_source or 'unknown'}")
                            debug_parts.append(f"publish_date={publish_date.isoformat()}")
                            if play_count:
                                debug_parts.append(f"play_count={play_count}")
                            self.log(f"[{username}] 日期诊断 | " + " | ".join(debug_parts))

                            if publish_date < start_d or publish_date > end_d:
                                continue

                            all_records.append({
                                "group_name": account_to_group.get(username, "未分组"),
                                "username": username,
                                "publish_date": publish_date.isoformat(),
                                "video_id": video_id,
                                "play_count": play_count,
                                "video_url": link,
                                "source_tab": url,
                            })
                            matched_count += 1

                        self.log(f"[{username}] 显式时间命中数：{has_explicit_time_count}")
                        self.log(f"[{username}] 反推时间命中数：{has_inferred_time_count}")
                        if has_detail_time_count:
                            self.log(f"[{username}] 详情页兜底命中数：{has_detail_time_count}")
                        self.log(f"[{username}] 日期范围内命中数量：{matched_count}")

                self.log(f"浏览器标签页总数：{page_count}")
                if skipped_tabs:
                    self.log(f"跳过的非账号主页标签页数：{len(skipped_tabs)}")

            finally:
                try:
                    browser.close()
                except Exception:
                    pass

        deduped = self.deduplicate_records(all_records)
        self.log(f"去重前：{len(all_records)} 条，去重后：{len(deduped)} 条")
        self.log(
            f"抓取诊断汇总：账号主页={len(discovered_accounts)} | "
            f"worker={min(self.max_workers, len(discovered_accounts) or 1)} | 明细={len(deduped)}"
        )

        deduped.sort(
            key=lambda x: (
                x["username"],
                x["publish_date"],
                -PUBLISH_CONFIDENCE_RANK.get((x.get("publish_confidence") or "").strip(), 0),
                x["video_url"],
            )
        )
        summary_order = target_accounts if target_accounts else sorted(set(discovered_accounts))
        summary_rows = self.build_summary(deduped, account_groups, summary_order)
        group_summary_rows = self.build_group_summary(deduped, account_groups)

        return deduped, summary_rows, group_summary_rows

    def collect_from_open_tabs_api_first(
            self,
            debug_url: str,
            start_d: date,
            end_d: date,
            target_tz,
            target_accounts=None,
            account_to_group=None,
            account_groups=None,
    ):
        target_accounts = target_accounts or []
        account_to_group = account_to_group or {}
        account_groups = account_groups or []

        all_records = []
        discovered_accounts = []
        account_diagnostics = []

        # 先让 GUI 显示“初始化中”
        self.report_progress(0, 0, "", "正在连接浏览器并扫描账号主页")

        with sync_playwright() as p:
            self.log(f"连接浏览器：{debug_url}")
            browser = p.chromium.connect_over_cdp(debug_url)

            try:
                self.ensure_target_account_pages_open(browser, target_accounts, refresh_existing=self.refresh_existing)
                profile_targets, page_count = self.collect_profile_targets(browser, target_accounts)
                discovered_accounts = [item["username"] for item in profile_targets]

                self.log(f"浏览器标签页总数：{page_count}")
            finally:
                try:
                    browser.close()
                except Exception:
                    pass

        if self.should_stop():
            self.log("检测到停止信号，终止任务。")
            return [], [], [], []

        if not profile_targets:
            self.log("未找到可抓取的账号主页标签页。")
            self.report_progress(0, 1, "", "未找到可抓取账号")
            summary_rows = self.build_summary([], account_groups, target_accounts)
            group_summary_rows = self.build_group_summary([], account_groups)
            return [], summary_rows, group_summary_rows, []

        total_accounts = len(profile_targets)
        self.report_progress(0, total_accounts, "", f"已获取待抓取账号，共 {total_accounts} 个")

        worker_count = min(self.max_workers, len(profile_targets))
        self.log(f"并发抓取账号主页：{len(profile_targets)} 个 | worker={worker_count}")

        if worker_count <= 1:
            for idx, item in enumerate(profile_targets, start=1):
                if self.should_stop():
                    self.log("检测到停止信号，终止任务。")
                    return [], [], [], []

                self.report_progress(idx - 1, total_accounts, item["username"], "开始抓取")

                result, diagnostic = self.collect_single_profile_api_first(
                    debug_url=debug_url,
                    username=item["username"],
                    url=item["url"],
                    start_d=start_d,
                    end_d=end_d,
                    target_tz=target_tz,
                    account_to_group=account_to_group,
                )
                all_records.extend(result)
                account_diagnostics.append(diagnostic)

                self.report_progress(idx, total_accounts, item["username"], "抓取完成")

        else:
            # 并发时按“完成账号数”推进
            completed_count = 0
            with ThreadPoolExecutor(max_workers=worker_count, thread_name_prefix="tiktok-crawl") as executor:
                future_to_item = {
                    executor.submit(
                        self.collect_single_profile_api_first,
                        debug_url,
                        item["username"],
                        item["url"],
                        start_d,
                        end_d,
                        target_tz,
                        account_to_group,
                    ): item
                    for item in profile_targets
                }

                for future in as_completed(future_to_item):
                    item = future_to_item[future]

                    if self.should_stop():
                        self.log("检测到停止信号，等待并发任务收尾。")

                    result, diagnostic = future.result()
                    all_records.extend(result)
                    account_diagnostics.append(diagnostic)

                    completed_count += 1
                    self.report_progress(completed_count, total_accounts, item["username"], "抓取完成")

        deduped = self.deduplicate_records(all_records)
        self.log(f"去重前：{len(all_records)} 条，去重后：{len(deduped)} 条")

        deduped.sort(key=lambda x: (x["username"], x["publish_date"], x["video_url"]))
        summary_order = target_accounts if target_accounts else sorted(set(discovered_accounts))
        summary_rows = self.build_summary(deduped, account_groups, summary_order)
        group_summary_rows = self.build_group_summary(deduped, account_groups)

        self.report_progress(total_accounts, total_accounts, "", "全部账号抓取完成")
        account_diagnostics.sort(key=lambda x: (x.get("group_name", ""), x.get("username", "")))
        return deduped, summary_rows, group_summary_rows, account_diagnostics

# =========================================================
# GUI
# =========================================================

class TikTokOpenTabsGUI:
    def __init__(self, root):
        self.root = root
        self.root.title("TikTok 已打开标签页抓取工具（最终版）")
        self.root.geometry("920x860")
        self.root.minsize(840, 760)

        self.stop_flag = False
        self.worker_thread = None
        self.account_groups = []
        self.group_name_var = tk.StringVar(value="")
        self.last_output_path = ""
        self.last_log_path = ""
        self.log_writer = None
        self.browser_ready_hint_var = tk.StringVar(
            value="第1步：启动专用浏览器；第2步：登录 TikTok 并打开账号主页；第3步：回来点开始抓取。"
        )
        self.progress_var = tk.DoubleVar(value=0.0)
        self.progress_text_var = tk.StringVar(value="未开始")
        self.progress_percent_var = tk.StringVar(value="0%")
        self.create_widgets()
        self.apply_compact_layout()
        self.load_ui_settings_from_disk()
        self.load_account_groups_from_disk()
        self.root.protocol("WM_DELETE_WINDOW", self.on_close)

    def create_widgets(self):
        frm_guide = ttk.LabelFrame(self.root, text="新手使用说明", padding=12)
        frm_guide.pack(fill="x", padx=12, pady=(12, 0))
        ttk.Label(
            frm_guide,
            text="1. 点击“启动专用浏览器”\n2. 在打开的浏览器里登录 TikTok，并打开一个或多个账号主页\n3. 回到本工具点击“开始抓取”",
            justify="left",
        ).pack(side="left", anchor="w")
        ttk.Label(frm_guide, textvariable=self.browser_ready_hint_var, foreground="#1f6feb").pack(
            side="right", anchor="e", padx=(12, 0)
        )
        guide_actions = ttk.Frame(frm_guide)
        guide_actions.pack(side="right", anchor="ne", padx=(12, 0))

        frm_top = ttk.Frame(self.root, padding=12)
        frm_top.pack(fill="x")

        self.debug_url_var = tk.StringVar(value="http://127.0.0.1:9222")
        self.refresh_existing_pages_var = tk.BooleanVar(value=True)
        self.allow_detail_fallback_var = tk.BooleanVar(value=True)
        self.debug_url_entry = ttk.Entry(frm_top, textvariable=self.debug_url_var, width=38)
        self.debug_url_entry.grid(row=0, column=1, sticky="w", padx=4, pady=4)
        self.btn_launch_browser = ttk.Button(guide_actions, text="启动专用浏览器", command=self.launch_browser_for_beginner)
        self.btn_launch_browser.pack(fill="x")
        try:
            self.debug_url_entry.grid_remove()
        except Exception:
            pass

        ttk.Label(frm_top, text="开始日期 (YYYY-MM-DD)：").grid(row=0, column=0, sticky="w", padx=4, pady=4)
        self.start_date_var = tk.StringVar(value="2026-03-01")
        ttk.Entry(frm_top, textvariable=self.start_date_var, width=16).grid(row=0, column=1, sticky="w", padx=4, pady=4)

        ttk.Label(frm_top, text="结束日期 (YYYY-MM-DD)：").grid(row=0, column=2, sticky="w", padx=18, pady=4)
        self.end_date_var = tk.StringVar(value="2026-03-31")
        ttk.Entry(frm_top, textvariable=self.end_date_var, width=16).grid(row=0, column=3, sticky="w", padx=4, pady=4)

        ttk.Label(frm_top, text="国家/时区：").grid(row=0, column=4, sticky="w", padx=18, pady=4)
        self.timezone_var = tk.StringVar(value="UTC+08:00")
        ttk.Combobox(
            frm_top,
            textvariable=self.timezone_var,
            values=list(TIMEZONE_OPTIONS.keys()),
            width=24,
            state="readonly",
        ).grid(row=0, column=5, sticky="w", padx=4, pady=4)

        ttk.Label(frm_top, text="Excel输出路径：").grid(row=1, column=0, sticky="w", padx=4, pady=4)
        self.output_var = tk.StringVar(value=os.path.abspath("tiktok_open_tabs_results.xlsx"))
        ttk.Entry(frm_top, textvariable=self.output_var, width=86).grid(row=1, column=1, columnspan=6, sticky="we", padx=4, pady=4)
        ttk.Checkbutton(frm_top, text="启动时刷新已打开主页", variable=self.refresh_existing_pages_var).grid(row=2, column=1, sticky="w", padx=4, pady=4)
        ttk.Checkbutton(frm_top, text="允许详情页兜底", variable=self.allow_detail_fallback_var).grid(row=2, column=2, sticky="w", padx=4, pady=4)

        ttk.Button(frm_top, text="选择路径", command=self.choose_output_file).grid(row=1, column=7, sticky="w", padx=4, pady=4)
        self.btn_open_output = ttk.Button(frm_top, text="打开文件", command=self.open_last_output_file, state="disabled")
        self.btn_open_output.grid(row=1, column=8, sticky="w", padx=4, pady=4)

        frm_accounts = ttk.LabelFrame(
            self.root,
            text="账号分组输入｜每次输入一批账号算一组，可自定义组名",
            padding=12,
        )
        frm_accounts.pack(fill="both", expand=False, padx=12, pady=8)

        frm_group_meta = ttk.Frame(frm_accounts)
        frm_group_meta.pack(fill="x", pady=(0, 8))

        ttk.Label(frm_group_meta, text="组名：").pack(side="left")
        ttk.Entry(frm_group_meta, textvariable=self.group_name_var, width=28).pack(side="left", padx=(4, 12))
        ttk.Button(frm_group_meta, text="保存/更新当前组", command=self.save_current_group).pack(side="left", padx=4)
        ttk.Button(frm_group_meta, text="加载选中组", command=self.load_selected_group).pack(side="left", padx=4)
        ttk.Button(frm_group_meta, text="删除选中组", command=self.delete_selected_group).pack(side="left", padx=4)
        ttk.Button(frm_group_meta, text="清空编辑区", command=self.clear_group_editor).pack(side="left", padx=4)

        self.accounts_text = tk.Text(frm_accounts, height=8, wrap="none")
        self.accounts_text.pack(fill="both", expand=True)

        demo_text = (
            "example_account_1\n"
            "@example_account_2\n"
            "https://www.tiktok.com/@example_account_3\n"
        )
        self.accounts_text.insert("1.0", demo_text)

        frm_group_list = ttk.Frame(frm_accounts)
        frm_group_list.pack(fill="both", expand=True, pady=(8, 0))

        self.groups_tree = ttk.Treeview(
            frm_group_list,
            columns=("group_name", "account_count", "accounts_preview"),
            show="headings",
            height=6,
            selectmode="extended",
        )
        self.groups_tree.heading("group_name", text="组名")
        self.groups_tree.heading("account_count", text="账号数")
        self.groups_tree.heading("accounts_preview", text="账号预览")
        self.groups_tree.column("group_name", width=180, anchor="w")
        self.groups_tree.column("account_count", width=80, anchor="center")
        self.groups_tree.column("accounts_preview", width=720, anchor="w")
        self.groups_tree.pack(side="left", fill="both", expand=True)

        scroll_groups = ttk.Scrollbar(frm_group_list, orient="vertical", command=self.groups_tree.yview)
        scroll_groups.pack(side="right", fill="y")
        self.groups_tree.configure(yscrollcommand=scroll_groups.set)

        self.btn_start = ttk.Button(guide_actions, text="开始抓取", command=self.start_task)
        self.btn_start.pack(fill="x", pady=(8, 0))

        self.btn_stop = ttk.Button(guide_actions, text="停止任务", command=self.stop_task)
        self.btn_stop.pack(fill="x", pady=(8, 0))

        frm_ops = ttk.Frame(self.root, padding=12)
        frm_ops.pack(fill="x")

        self.btn_clear = ttk.Button(frm_ops, text="清空日志", command=self.clear_log)
        self.btn_clear.pack(side="left", padx=6)

        self.btn_check_browser = ttk.Button(frm_ops, text="检测浏览器连接", command=self.check_browser_connection)
        self.btn_check_browser.pack(side="left", padx=6)

        self.btn_list_tabs = ttk.Button(frm_ops, text="列出已打开标签页", command=self.list_open_tabs)
        self.btn_list_tabs.pack(side="left", padx=6)

        frm_progress = ttk.LabelFrame(self.root, text="任务进度", padding=12)
        frm_progress.pack(fill="x", padx=12, pady=(0, 8))

        progress_top = ttk.Frame(frm_progress)
        progress_top.pack(fill="x")

        ttk.Label(progress_top, textvariable=self.progress_text_var).pack(side="left", anchor="w")
        ttk.Label(progress_top, textvariable=self.progress_percent_var).pack(side="right", anchor="e")

        self.progress_bar = ttk.Progressbar(
            frm_progress,
            variable=self.progress_var,
            maximum=100,
            mode="determinate",
        )
        self.progress_bar.pack(fill="x", pady=(8, 0))
        frm_log = ttk.LabelFrame(self.root, text="运行日志", padding=12)
        frm_log.pack(fill="both", expand=True, padx=12, pady=8)

        self.log_text = tk.Text(frm_log, height=28, wrap="word")
        self.log_text.pack(side="left", fill="both", expand=True)

        scroll = ttk.Scrollbar(frm_log, orient="vertical", command=self.log_text.yview)
        scroll.pack(side="right", fill="y")
        self.log_text.configure(yscrollcommand=scroll.set)

    def apply_compact_layout(self):
        root_children = self.root.winfo_children()
        if len(root_children) < 5:
            return

        frm_guide = root_children[0]
        frm_top = root_children[1]
        frm_accounts = root_children[2]
        frm_ops = root_children[3]

        try:
            guide_labels = frm_guide.winfo_children()
            if len(guide_labels) >= 2:
                guide_labels[0].configure(wraplength=520, justify="left")
                guide_labels[1].configure(wraplength=520, justify="left")
        except Exception:
            pass

        top_widgets = frm_top.winfo_children()
        if len(top_widgets) >= 13:
            for widget in top_widgets:
                try:
                    widget.grid_forget()
                except Exception:
                    pass

            for col in range(6):
                frm_top.columnconfigure(col, weight=1 if col in (1, 3, 5) else 0)

            # 顶部区域按三组标签加输入框重排，避免紧凑布局漏掉日期控件。
            top_widgets[1].grid(row=0, column=0, sticky="w", padx=4, pady=4)
            top_widgets[2].grid(row=0, column=1, sticky="we", padx=4, pady=4)
            top_widgets[3].grid(row=0, column=2, sticky="w", padx=18, pady=4)
            top_widgets[4].grid(row=0, column=3, sticky="we", padx=4, pady=4)
            top_widgets[5].grid(row=0, column=4, sticky="w", padx=18, pady=4)
            top_widgets[6].grid(row=0, column=5, sticky="we", padx=4, pady=4)

            top_widgets[7].grid(row=1, column=0, sticky="w", padx=4, pady=4)
            top_widgets[8].grid(row=1, column=1, columnspan=3, sticky="we", padx=4, pady=4)
            top_widgets[11].grid(row=1, column=4, sticky="we", padx=4, pady=4)
            top_widgets[12].grid(row=1, column=5, sticky="we", padx=4, pady=4)

            top_widgets[9].grid(row=2, column=1, columnspan=2, sticky="w", padx=4, pady=(2, 4))
            top_widgets[10].grid(row=2, column=3, columnspan=3, sticky="w", padx=4, pady=(2, 4))

        account_children = frm_accounts.winfo_children()
        if len(account_children) >= 3:
            frm_group_meta = account_children[0]
            accounts_text = account_children[1]
            group_list_frame = account_children[2]

            meta_widgets = frm_group_meta.winfo_children()
            if len(meta_widgets) >= 6:
                for widget in meta_widgets:
                    try:
                        widget.pack_forget()
                    except Exception:
                        pass
                frm_group_meta.columnconfigure(1, weight=1)
                meta_widgets[0].grid(row=0, column=0, sticky="w")
                meta_widgets[1].grid(row=0, column=1, columnspan=3, sticky="we", padx=(6, 0))
                meta_widgets[2].grid(row=1, column=0, sticky="w", pady=(8, 0))
                meta_widgets[3].grid(row=1, column=1, sticky="w", padx=(6, 0), pady=(8, 0))
                meta_widgets[4].grid(row=1, column=2, sticky="w", padx=(6, 0), pady=(8, 0))
                meta_widgets[5].grid(row=1, column=3, sticky="w", padx=(6, 0), pady=(8, 0))

            try:
                accounts_text.configure(height=6)
            except Exception:
                pass

            try:
                tree = group_list_frame.winfo_children()[0]
                tree.configure(height=5)
                tree.column("group_name", width=150, anchor="w")
                tree.column("account_count", width=70, anchor="center")
                tree.column("accounts_preview", width=480, anchor="w")
            except Exception:
                pass

        op_widgets = frm_ops.winfo_children()
        if op_widgets:
            row1 = ttk.Frame(frm_ops)
            row1.pack(fill="x")
            for widget in op_widgets:
                try:
                    widget.pack_forget()
                except Exception:
                    pass
            for widget in op_widgets:
                widget.pack(in_=row1, side="left", padx=4)

        try:
            self.log_text.configure(height=22)
        except Exception:
            pass

    def _normalize_group_name(self, name: str) -> str:
        return (name or "").strip() or "未命名分组"

    def _get_editor_accounts(self):
        raw_accounts = self.accounts_text.get("1.0", "end").strip()
        if not raw_accounts:
            return []
        return parse_accounts_from_text(raw_accounts)

    def _get_selected_group_index(self):
        sels = self.groups_tree.selection()
        if not sels:
            return None
        try:
            return int(str(sels[0]).replace("group_", ""))
        except Exception:
            return None

    def _get_selected_group_indices(self):
        indices = []
        for sel in self.groups_tree.selection():
            try:
                idx = int(str(sel).replace("group_", ""))
            except Exception:
                continue
            if 0 <= idx < len(self.account_groups):
                indices.append(idx)
        return sorted(set(indices))

    def refresh_groups_tree(self):
        for item in self.groups_tree.get_children():
            self.groups_tree.delete(item)

        for idx, group in enumerate(self.account_groups):
            accounts = list(group.get("accounts", []))
            preview = ", ".join(accounts[:6])
            if len(accounts) > 6:
                preview += f" ... 共 {len(accounts)} 个"
            self.groups_tree.insert(
                "",
                "end",
                iid=f"group_{idx}",
                values=(
                    group.get("group_name", ""),
                    len(accounts),
                    preview,
                ),
            )

    def _normalized_account_groups_payload(self):
        out = []
        seen_groups = set()
        for group in self.account_groups:
            if not isinstance(group, dict):
                continue
            group_name = self._normalize_group_name(group.get("group_name", ""))
            if group_name in seen_groups:
                continue
            seen_groups.add(group_name)
            accounts = []
            seen_accounts = set()
            for account in group.get("accounts", []) or []:
                try:
                    username = normalize_username(account)
                except Exception:
                    continue
                if not username or username in seen_accounts:
                    continue
                seen_accounts.add(username)
                accounts.append(username)
            if accounts:
                out.append({
                    "group_name": group_name,
                    "accounts": accounts,
                })
        return out

    def save_account_groups_to_disk(self):
        payload = {
            "version": 1,
            "account_groups": self._normalized_account_groups_payload(),
        }
        with open(ACCOUNT_GROUPS_FILE, "w", encoding="utf-8") as f:
            json.dump(payload, f, ensure_ascii=False, indent=2)

    def load_account_groups_from_disk(self):
        if not os.path.exists(ACCOUNT_GROUPS_FILE):
            return
        try:
            with open(ACCOUNT_GROUPS_FILE, "r", encoding="utf-8") as f:
                data = json.load(f)
            groups = data.get("account_groups", []) if isinstance(data, dict) else []
            if not isinstance(groups, list):
                groups = []
            self.account_groups = []
            for group in groups:
                if not isinstance(group, dict):
                    continue
                group_name = self._normalize_group_name(group.get("group_name", ""))
                accounts = []
                seen = set()
                for account in group.get("accounts", []) or []:
                    try:
                        username = normalize_username(account)
                    except Exception:
                        continue
                    if not username or username in seen:
                        continue
                    seen.add(username)
                    accounts.append(username)
                if accounts:
                    self.account_groups.append({
                        "group_name": group_name,
                        "accounts": accounts,
                    })
            self.refresh_groups_tree()
            self.log(f"已恢复本地账号分组：{len(self.account_groups)} 组")
        except Exception as e:
            self.log(f"读取本地账号分组失败：{e}")

    def save_ui_settings_to_disk(self):
        payload = {
            "window_geometry": self.root.winfo_geometry(),
            "start_date": (self.start_date_var.get() or "").strip(),
            "end_date": (self.end_date_var.get() or "").strip(),
            "timezone": normalize_timezone_name(self.timezone_var.get()),
            "output_path": (self.output_var.get() or "").strip(),
            "debug_url": (self.debug_url_var.get() or "").strip(),
            "refresh_existing_pages": bool(self.refresh_existing_pages_var.get()),
            "allow_detail_fallback": bool(self.allow_detail_fallback_var.get()),
        }
        with open(UI_SETTINGS_FILE, "w", encoding="utf-8") as f:
            json.dump(payload, f, ensure_ascii=False, indent=2)

    def load_ui_settings_from_disk(self):
        if not os.path.exists(UI_SETTINGS_FILE):
            return
        try:
            with open(UI_SETTINGS_FILE, "r", encoding="utf-8") as f:
                data = json.load(f)
            if not isinstance(data, dict):
                return
            geometry = (data.get("window_geometry") or "").strip()
            if geometry:
                try:
                    self.root.geometry(geometry)
                except Exception:
                    pass
            if data.get("start_date"):
                self.start_date_var.set(str(data.get("start_date")))
            if data.get("end_date"):
                self.end_date_var.set(str(data.get("end_date")))
            if data.get("timezone"):
                self.timezone_var.set(normalize_timezone_name(str(data.get("timezone"))))
            if data.get("output_path"):
                self.output_var.set(str(data.get("output_path")))
            if data.get("debug_url"):
                self.debug_url_var.set(str(data.get("debug_url")))
            if "refresh_existing_pages" in data:
                self.refresh_existing_pages_var.set(bool(data.get("refresh_existing_pages")))
            if "allow_detail_fallback" in data:
                self.allow_detail_fallback_var.set(bool(data.get("allow_detail_fallback")))
        except Exception as e:
            self.log(f"读取界面设置失败：{e}")

    def on_close(self):
        try:
            self.save_ui_settings_to_disk()
        except Exception as e:
            try:
                self.log(f"保存界面设置失败：{e}")
            except Exception:
                pass
        self.root.destroy()

    def save_current_group(self):
        try:
            accounts = self._get_editor_accounts()
        except Exception as e:
            messagebox.showerror("账号格式错误", str(e))
            return

        if not accounts:
            messagebox.showwarning("提示", "请先输入当前组的账号列表。")
            return

        group_name = self._normalize_group_name(self.group_name_var.get())
        payload = {"group_name": group_name, "accounts": accounts}

        replaced = False
        for idx, group in enumerate(self.account_groups):
            if group.get("group_name") == group_name:
                self.account_groups[idx] = payload
                replaced = True
                break

        if not replaced:
            self.account_groups.append(payload)

        self.account_groups = self._normalized_account_groups_payload()
        self.refresh_groups_tree()
        self.save_account_groups_to_disk()
        self.log(f"已保存分组：{group_name} | 账号数：{len(accounts)}")

    def load_selected_group(self):
        idx = self._get_selected_group_index()
        if idx is None or idx >= len(self.account_groups):
            messagebox.showwarning("提示", "请先选中一个分组。")
            return

        group = self.account_groups[idx]
        self.group_name_var.set(group.get("group_name", ""))
        self.accounts_text.delete("1.0", "end")
        self.accounts_text.insert("1.0", "\n".join(group.get("accounts", [])))
        self.log(f"已加载分组到编辑区：{group.get('group_name', '')}")

    def delete_selected_group(self):
        idx = self._get_selected_group_index()
        if idx is None or idx >= len(self.account_groups):
            messagebox.showwarning("提示", "请先选中一个分组。")
            return

        group_name = self.account_groups[idx].get("group_name", "")
        del self.account_groups[idx]
        self.account_groups = self._normalized_account_groups_payload()
        self.refresh_groups_tree()
        self.save_account_groups_to_disk()
        self.log(f"已删除分组：{group_name}")

    def clear_group_editor(self):
        self.group_name_var.set("")
        self.accounts_text.delete("1.0", "end")

    def build_groups_for_run(self):
        selected_indices = self._get_selected_group_indices()
        groups = []
        for idx in selected_indices:
            group = self.account_groups[idx]
            if not group.get("accounts"):
                continue
            groups.append({
                "group_name": self._normalize_group_name(group.get("group_name", "")),
                "accounts": list(group.get("accounts", [])),
            })

        editor_accounts = self._get_editor_accounts()
        if editor_accounts:
            editor_group_name = self._normalize_group_name(self.group_name_var.get())
            merged = False
            for group in groups:
                if group["group_name"] == editor_group_name:
                    group["accounts"] = editor_accounts
                    merged = True
                    break
            if not merged and not selected_indices:
                groups.append({
                    "group_name": editor_group_name,
                    "accounts": editor_accounts,
                })

        return groups

    def choose_output_file(self):
        path = filedialog.asksaveasfilename(
            title="选择导出 Excel 文件",
            defaultextension=".xlsx",
            filetypes=[("Excel 文件", "*.xlsx")],
        )
        if path:
            self.output_var.set(path)

    @staticmethod
    def build_timestamped_output_path(base_path: str) -> str:
        raw = (base_path or "").strip()
        if not raw:
            return ""
        if not raw.lower().endswith(".xlsx"):
            raw += ".xlsx"
        abs_path = os.path.abspath(raw)
        folder = os.path.dirname(abs_path) or os.getcwd()
        filename = os.path.basename(abs_path)
        stem, ext = os.path.splitext(filename)
        stem = re.sub(r"_\d{8}_\d{6}$", "", stem)
        ts = datetime.now().strftime("%Y%m%d_%H%M%S")
        return os.path.join(folder, f"{stem}_{ts}{ext or '.xlsx'}")

    def open_last_output_file(self):
        path = (self.last_output_path or "").strip()
        if not path:
            messagebox.showinfo("提示", "当前还没有已导出的文件。")
            return
        if not os.path.exists(path):
            messagebox.showwarning("提示", f"文件不存在：\n{path}")
            return
        try:
            os.startfile(path)
        except Exception as e:
            messagebox.showerror("错误", f"打开文件失败：{e}")

    def log(self, msg: str):
        try:
            if self.log_writer is not None:
                self.log_writer.write(msg)
        except Exception:
            pass
        self.root.after(0, self._append_log, msg)

    def _append_log(self, msg: str):
        now = datetime.now().strftime("%H:%M:%S")
        self.log_text.insert("end", f"[{now}] {msg}\n")
        self.log_text.see("end")

    def reset_progress(self):
        self.root.after(0, self._reset_progress_ui)

    def _reset_progress_ui(self):
        try:
            self.progress_bar.stop()
        except Exception:
            pass
        self.progress_bar.configure(mode="indeterminate", maximum=100)
        self.progress_var.set(0)
        self.progress_text_var.set("准备开始...")
        self.progress_percent_var.set("0%")
        try:
            self.progress_bar.start(10)
        except Exception:
            pass

    def update_progress(self, current: int, total: int, username: str = "", phase: str = ""):
        self.root.after(0, self._update_progress_ui, current, total, username, phase)

    def _update_progress_ui(self, current: int, total: int, username: str = "", phase: str = ""):
        try:
            self.progress_bar.stop()
        except Exception:
            pass

        # total <= 0 时，显示不确定进度
        if total <= 0:
            self.progress_bar.configure(mode="indeterminate", maximum=100)
            try:
                self.progress_bar.start(10)
            except Exception:
                pass
            text = phase or "处理中..."
            if username:
                text = f"{text}  @{username}"
            self.progress_text_var.set(text)
            self.progress_percent_var.set("...")
            return

        safe_total = max(1, int(total))
        safe_current = max(0, min(int(current), safe_total))

        self.progress_bar.configure(mode="determinate", maximum=safe_total)
        self.progress_var.set(safe_current)

        percent = (safe_current / safe_total) * 100
        self.progress_percent_var.set(f"{percent:.0f}%")

        text = f"{safe_current}/{safe_total}"
        if phase:
            text += f" | {phase}"
        if username:
            text += f" | @{username}"

        self.progress_text_var.set(text)

    def finish_progress(self, success: bool = True, text: str = ""):
        self.root.after(0, self._finish_progress_ui, success, text)

    def _finish_progress_ui(self, success: bool = True, text: str = ""):
        try:
            self.progress_bar.stop()
        except Exception:
            pass

        self.progress_bar.configure(mode="determinate")

        current_max = self.progress_bar.cget("maximum")
        try:
            current_max = float(current_max)
        except Exception:
            current_max = 100.0

        if success:
            self.progress_var.set(current_max)
            self.progress_percent_var.set("100%")
            self.progress_text_var.set(text or "任务完成")
        else:
            self.progress_text_var.set(text or "任务已停止/失败")

    def clear_log(self):
        self.log_text.delete("1.0", "end")

    def stop_task(self):
        self.stop_flag = True
        self.log("已发送停止信号，当前步骤结束后会自动停止")
        self.update_progress(0, 0, "", "正在停止，请稍候")

    def launch_browser_for_beginner(self):
        debug_url = self.debug_url_var.get().strip() or "http://127.0.0.1:9222"
        ok, path_or_err = launch_debug_browser(debug_url, log_func=self.log)
        if ok:
            self.browser_ready_hint_var.set("浏览器已启动：请在新浏览器中登录 TikTok，并打开账号主页后再开始抓取。")
            messagebox.showinfo(
                "浏览器已启动",
                "已启动专用浏览器。\n\n请在浏览器中：\n1. 登录 TikTok\n2. 打开一个或多个账号主页\n3. 回到本工具点击“开始抓取”"
            )
        else:
            messagebox.showerror("启动失败", str(path_or_err))

    def ensure_browser_ready_for_beginner(self):
        debug_url = self.debug_url_var.get().strip() or "http://127.0.0.1:9222"
        ok, result = check_cdp_endpoint(debug_url)
        if not ok:
            self.log(f"未检测到可连接浏览器，准备自动启动：{result}")
            ok2, result2 = launch_debug_browser(debug_url, log_func=self.log)
            if not ok2:
                return False, result2
        self.browser_ready_hint_var.set("浏览器已就绪：请确认已经登录 TikTok 并打开账号主页。")
        return True, debug_url

    def check_browser_connection(self):
        debug_url = (self.debug_url_var.get() or "").strip()
        if not debug_url:
            messagebox.showwarning("提示", "请先填写调试浏览器地址")
            return
        ok, result = check_cdp_endpoint(debug_url)
        if ok:
            self.log(f"浏览器连接检测成功：{result}")
            messagebox.showinfo("检测成功", f"已成功连接调试浏览器：\n{result}")
        else:
            self.log(f"浏览器连接检测失败：{result}")
            messagebox.showerror("检测失败", str(result))

    def list_open_tabs(self):
        debug_url = (self.debug_url_var.get() or "").strip()
        if not debug_url:
            messagebox.showwarning("提示", "请先填写调试浏览器地址")
            return
        ok, result = check_cdp_endpoint(debug_url)
        if not ok:
            self.log(f"浏览器连接检测失败：{result}")
            messagebox.showerror("检测失败", str(result))
            return

        try:
            with sync_playwright() as p:
                browser = p.chromium.connect_over_cdp(debug_url)
                tabs = []
                try:
                    for context in browser.contexts:
                        for page in context.pages:
                            tabs.append({
                                "title": (page.title() or "").strip() or "<无标题>",
                                "url": (page.url or "").strip() or "<空地址>",
                            })
                finally:
                    try:
                        browser.close()
                    except Exception:
                        pass
        except Exception as e:
            self.log(f"列出标签页失败：{e}")
            messagebox.showerror("错误", f"列出标签页失败：{e}")
            return

        self.log("=" * 72)
        self.log(f"当前已打开标签页：{len(tabs)} 个")
        for idx, tab in enumerate(tabs, start=1):
            self.log(f"[{idx:02d}] {tab['title']}")
            self.log(f"     {tab['url']}")

        if not tabs:
            messagebox.showinfo("结果", "当前未读取到任何已打开标签页。")
        else:
            messagebox.showinfo("结果", f"已读取到 {len(tabs)} 个已打开标签页，详情见日志。")

    def validate_inputs(self):
        debug_url = (self.debug_url_var.get() or "").strip()
        if not debug_url:
            raise ValueError("请填写调试浏览器地址，例如 http://127.0.0.1:9222")

        try:
            start_d = datetime.strptime((self.start_date_var.get() or "").strip(), "%Y-%m-%d").date()
        except Exception:
            raise ValueError("开始日期格式错误，请使用 YYYY-MM-DD")

        try:
            end_d = datetime.strptime((self.end_date_var.get() or "").strip(), "%Y-%m-%d").date()
        except Exception:
            raise ValueError("结束日期格式错误，请使用 YYYY-MM-DD")

        timezone_name = normalize_timezone_name(self.timezone_var.get())
        try:
            self.timezone_var.set(timezone_name)
        except Exception:
            pass
        target_tz = TIMEZONE_OPTIONS.get(timezone_name)
        if target_tz is None:
            raise ValueError("请选择有效的国家/时区")

        if start_d > end_d:
            raise ValueError("开始日期不能大于结束日期")

        account_groups = self.build_groups_for_run()
        accounts = []
        account_to_group = {}
        duplicates = []
        for group in account_groups:
            group_name = self._normalize_group_name(group.get("group_name", ""))
            for account in group.get("accounts", []):
                username = normalize_username(account)
                if username in account_to_group and account_to_group[username] != group_name:
                    duplicates.append(f"@{username}: {account_to_group[username]} / {group_name}")
                    continue
                account_to_group[username] = group_name
                if username not in accounts:
                    accounts.append(username)

        selected = self._get_selected_group_indices()
        if selected and not account_groups:
            raise ValueError("请先在组列表中选中至少一个分组，再开始抓取。")
        if duplicates:
            raise ValueError("以下账号被分配到了多个分组，请先修正：\n" + "\n".join(sorted(set(duplicates))))

        output_path = (self.output_var.get() or "").strip()
        if not output_path:
            raise ValueError("请设置导出 Excel 路径")
        if not output_path.lower().endswith(".xlsx"):
            output_path += ".xlsx"

        return debug_url, start_d, end_d, timezone_name, target_tz, account_groups, accounts, account_to_group, output_path

    def start_task(self):
        if self.worker_thread and self.worker_thread.is_alive():
            messagebox.showwarning("提示", "任务正在运行，请勿重复启动")
            return

        try:
            browser_ok, browser_result = self.ensure_browser_ready_for_beginner()
            if not browser_ok:
                raise ValueError(browser_result)
            debug_url, start_d, end_d, timezone_name, target_tz, account_groups, accounts, account_to_group, output_path = self.validate_inputs()
        except Exception as e:
            messagebox.showerror("参数错误", str(e))
            return

        output_path = self.build_timestamped_output_path(output_path)
        self.last_output_path = output_path
        try:
            self.btn_open_output.config(state="disabled")
        except Exception:
            pass

        self.stop_flag = False
        self.btn_start.config(state="disabled")
        self.reset_progress()
        self.log_writer = RunLogWriter(LOGS_DIR)
        self.last_log_path = self.log_writer.start()
        self.log("=" * 84)
        self.log("任务开始")
        self.log(f"运行日志：{self.last_log_path}")
        self.log(f"调试浏览器：{debug_url}")
        self.log(f"日期范围：{start_d} ~ {end_d}")
        self.log(f"国家/时区：{timezone_name}")
        self.log(f"Excel输出：{output_path}")

        if account_groups:
            self.log(f"分组模式：本次共 {len(account_groups)} 个分组，账号数 {len(accounts)}")
            for group in account_groups:
                self.log(f"  - {group['group_name']}：{len(group['accounts'])} 个账号")
        elif accounts:
            self.log(f"指定账号模式：本次仅统计这 {len(accounts)} 个账号")
        else:
            self.log("未填写账号分组，默认抓取当前已打开的所有 TikTok 账号主页")

        self.worker_thread = threading.Thread(
            target=self.run_crawler_task,
            args=(debug_url, start_d, end_d, target_tz, account_groups, accounts, account_to_group, output_path),
            daemon=True,
        )
        self.worker_thread.start()

    def run_crawler_task(self, debug_url, start_d, end_d, target_tz, account_groups, accounts, account_to_group,
                         output_path):
        try:
            crawler = TikTokOpenTabsCrawler(
                log_func=self.log,
                stop_flag_func=lambda: self.stop_flag,
                refresh_existing=bool(self.refresh_existing_pages_var.get()),
                allow_detail_fallback=bool(self.allow_detail_fallback_var.get()),
                progress_callback=self.update_progress,
            )

            records, summary_rows, group_summary_rows, diagnostic_rows = crawler.collect_from_open_tabs_api_first(
                debug_url=debug_url,
                start_d=start_d,
                end_d=end_d,
                target_tz=target_tz,
                target_accounts=accounts,
                account_to_group=account_to_group,
                account_groups=account_groups,
            )

            export_to_excel(
                records,
                summary_rows,
                group_summary_rows,
                output_path,
                diagnostic_rows=diagnostic_rows,
                log_path=self.last_log_path,
            )

            self.log("=" * 84)
            self.log(f"任务完成，明细记录数：{len(records)}")
            self.log(f"Excel 已导出：{output_path}")
            if self.last_log_path:
                self.log(f"日志已保存：{self.last_log_path}")
            self.last_output_path = output_path
            self.finish_progress(True, f"任务完成 | 明细 {len(records)} 条")

            self.root.after(
                0,
                lambda: (
                    self.btn_open_output.config(state="normal"),
                    messagebox.showinfo(
                        "完成",
                        f"抓取完成。\n\n明细记录数：{len(records)}\n导出文件：{output_path}\n日志文件：{self.last_log_path}",
                    ),
                ),
            )

        except Exception as e:
            err = f"运行失败：{e}\n\n{traceback.format_exc()}"
            self.log(err)
            self.finish_progress(False, "任务失败")
            err_msg = str(e)
            self.root.after(0, lambda msg=err_msg: messagebox.showerror("错误", msg))
        finally:
            if self.stop_flag:
                self.finish_progress(False, "任务已停止")
            self.log_writer = None
            self.root.after(0, lambda: self.btn_start.config(state="normal"))

def main():
    root = tk.Tk()
    app = TikTokOpenTabsGUI(root)
    root.mainloop()


if __name__ == "__main__":
    main()
