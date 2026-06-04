import tempfile
import unittest
from pathlib import Path

from check_empty_product_white_bg import export_to_excel, scan_white_bg_folders
from openpyxl import load_workbook


class WhiteBgAuditTests(unittest.TestCase):
    def test_scan_marks_empty_and_non_empty_target_folders(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            empty = root / "A" / "01.产品白底图"
            filled = root / "B" / "01.产品白底图"
            empty.mkdir(parents=True)
            filled.mkdir(parents=True)
            (filled / "photo.JPG").write_bytes(b"x")

            rows = scan_white_bg_folders(root)

            statuses = {Path(row.folder_path).parent.name: row.status for row in rows}
            self.assertEqual(statuses, {"A": "无图片", "B": "有图片"})

    def test_scan_ignores_images_inside_target_subfolders(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            target = root / "A" / "01.产品白底图"
            nested = target / "nested"
            nested.mkdir(parents=True)
            (nested / "photo.png").write_bytes(b"x")

            rows = scan_white_bg_folders(root)

            self.assertEqual(len(rows), 1)
            self.assertEqual(rows[0].status, "无图片")
            self.assertEqual(rows[0].image_count, 0)
            self.assertEqual(rows[0].image_files, [])

    def test_export_to_excel_creates_expected_content(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            target = root / "Product" / "01.产品白底图"
            target.mkdir(parents=True)
            output = root / "audit.xlsx"
            rows = scan_white_bg_folders(root)

            export_to_excel(rows, output)

            wb = load_workbook(output)
            ws = wb.active
            self.assertEqual(ws["A1"].value, "序号")
            self.assertEqual(ws["B1"].value, "状态")
            self.assertEqual(ws["B2"].value, "无图片")
            self.assertEqual(ws["D2"].value, str(target.parent))
            self.assertEqual(ws["E2"].value, 0)

    def test_scan_marks_product_folder_without_target_subfolder(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            product = root / "ProductWithoutWhiteBg"
            product.mkdir()

            rows = scan_white_bg_folders(root)

            self.assertEqual(len(rows), 1)
            self.assertEqual(rows[0].status, "无白底图")
            self.assertEqual(rows[0].folder_path, "")
            self.assertEqual(rows[0].product_dir, str(product))
            self.assertEqual(rows[0].image_count, 0)


if __name__ == "__main__":
    unittest.main()
